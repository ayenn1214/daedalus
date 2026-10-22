"""Document = ordered feature history + undo + rebuild + I/O. Qt-free."""
from __future__ import annotations

import json

import numpy as np
import re
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

from build123d import export_step, export_stl, import_step

from .features import FEATURE_TYPES, BuildContext, Feature, FeatureError
from .layers import DEFAULT, PALETTE, Layer, attribute_faces, valid_name
from . import assembly as asm
from .assembly import AssemblyError, Mate, Obj
from .physics import DRIVE_KINDS, Drive, PhysicsError, PhysicsSettings

from . import materials as matlib
from .materials import Material
from .doc_v28 import DEFAULT_CONFIG, V28Mixin

MATERIALS = {m.name: m.density for m in matlib.LIBRARY}   # name -> density g/cm³ (back-compat)


@dataclass
class RebuildResult:
    body: object | None = None
    sketches: dict = field(default_factory=dict)        # name -> (display shape, plane)
    sketch_status: dict = field(default_factory=dict)   # name -> {status, dof}
    sketch_issues: dict = field(default_factory=dict)   # name -> reason it has no profile
    cutlists: dict = field(default_factory=dict)        # frame feature -> members
    face_layers: list = field(default_factory=list)     # layer per face (single-object documents)
    bodies: dict = field(default_factory=dict)          # object -> body in its own coordinates
    face_layers_by_obj: dict = field(default_factory=dict)
    feature_obj: dict = field(default_factory=dict)     # feature name -> object
    face_owner_by_obj: dict = field(default_factory=dict)  # object -> feature that made each face
    assembly: object | None = None                      # assembly.SolveResult
    flats: dict = field(default_factory=dict)           # sheet-metal feature -> flat pattern
    threads: dict = field(default_factory=dict)         # cosmetic threads (by feature)
    planes: dict = field(default_factory=dict)          # reference plane feature -> (Plane, object)
    surfaces: dict = field(default_factory=dict)        # surface feature -> (shape, object)   (v2.8)
    errors: dict[str, str] = field(default_factory=dict)
    ms: float = 0.0


def _coerce_opt(default, v):
    try:
        if isinstance(default, bool):
            return v if isinstance(v, bool) else str(v).lower() in ("1", "true", "yes", "on")
        if isinstance(default, int):
            return int(float(v))
        if isinstance(default, float):
            return float(v)
    except (TypeError, ValueError) as e:
        raise FeatureError(f"Bad value {v!r}") from e
    return "" if v is None else str(v)


def _strip_cache(o):
    """Drop '_'-prefixed cache keys (computed view geometry) recursively."""
    if isinstance(o, dict):
        return {k: _strip_cache(v) for k, v in o.items() if not str(k).startswith("_")}
    if isinstance(o, list):
        return [_strip_cache(v) for v in o]
    return o


_VALID: dict = {}


def _valid_cached(body) -> bool:
    """body.is_valid, checked solid by solid and remembered per solid — after an edit only the solids that
    actually changed are checked again (1 500 placed copies of one imported screw are checked once). OCCT hashes
    shapes by TShape + location, so a changed solid never reuses an old answer. (is_valid is a property in some
    build123d versions and a method in others.)"""
    def check(shape):
        v = getattr(shape, "is_valid", True)
        return bool(v() if callable(v) else v)
    try:
        from OCP.TopAbs import TopAbs_SOLID
        from OCP.TopExp import TopExp_Explorer
        from build123d import Solid
        ex = TopExp_Explorer(body.wrapped, TopAbs_SOLID)
        sols = []
        while ex.More():
            sols.append(ex.Current())
            ex.Next()
    except Exception:  # noqa: BLE001
        return check(body)
    if not sols:
        return check(body)
    for sh in sols:
        try:
            key = hash(sh)
        except Exception:  # noqa: BLE001
            return check(body)
        hit = _VALID.get(key)
        if hit is None:
            hit = _VALID[key] = check(Solid(sh))
            if len(_VALID) > 20000:
                _VALID.clear()
        if not hit:
            return False
    return True


def nfaces(body) -> int:
    """Face count, remembered on the body (bodies are reused while their history is unchanged)."""
    n = getattr(body, "_dd_nf", None)
    if n is None:
        try:
            from OCP.TopAbs import TopAbs_FACE
            from OCP.TopExp import TopExp
            from OCP.TopTools import TopTools_IndexedMapOfShape
            m = TopTools_IndexedMapOfShape()
            (getattr(TopExp, "MapShapes_s", None) or TopExp.MapShapes)(body.wrapped, TopAbs_FACE, m)
            n = m.Extent()
        except Exception:  # noqa: BLE001
            n = len(body.faces())
        try:
            body._dd_nf = n
        except Exception:  # noqa: BLE001
            pass
    return n


class _LazyMap(dict):
    """dict whose values are computed on first access from per-key makers."""

    def __init__(self, makers: dict):
        super().__init__()
        self._makers = makers

    def __missing__(self, key):
        mk = self._makers.get(key)
        if mk is None:
            raise KeyError(key)
        v = self[key] = mk()
        return v

    def get(self, key, default=None):
        try:
            return self[key]
        except KeyError:
            return default

    def __contains__(self, key):
        return dict.__contains__(self, key) or key in self._makers


def _owners_cached(body, tools, names):
    """attribute_faces, memoised on the body object (bodies are reused while their history is unchanged)."""
    key = tuple(names)
    memo = getattr(body, "_dd_owners", None)
    if memo is None:
        memo = {}
        try:
            body._dd_owners = memo
        except Exception:  # noqa: BLE001
            pass
    if key not in memo:
        memo[key] = attribute_faces(body, tools, list(names))
    return memo[key]


class Document(V28Mixin):
    def __init__(self):
        self.features: list[Feature] = []
        self.rollback: int | None = None      # evaluate features[:rollback]
        self.material = matlib.DEFAULT
        self.custom_materials: dict[str, dict] = {}
        self.path: Path | None = None
        self.imported_step: str | None = None
        self.result = RebuildResult()
        self.dirty = False
        self._undo: list[str] = []
        self._redo: list[str] = []
        self._listeners: list[Callable[[], None]] = []
        self.layers: list[Layer] = [Layer(DEFAULT)]
        self.active_layer = DEFAULT
        self.objects: list[Obj] = [Obj("Object1", fixed=True)]
        self.active_object = "Object1"
        self.mates: list[Mate] = []
        self.physics = PhysicsSettings()
        self.last_sim = None                  # physics.SimResult of the latest run (not saved)
        self.studies: list[dict] = []         # FEA / CFD studies (saved)
        self.explode: dict = {}               # object -> exploded-view offset (saved)
        self.exploded = False
        self.sim_results: dict = {}           # study name -> result (not saved)
        self.drawings: list = []              # daedalus.drawing.model.Drawing (saved, caches stripped)
        self.export_options: dict = {}        # overrides of io.formats.EXPORT_DEFAULTS (saved)
        self.cam: dict = {}                   # CAM setup: stock, tools, operations (saved)
        self._v28_init()                      # equations, configurations, subassemblies, component patterns

    # ------------------------------------------------------------ events
    def subscribe(self, fn: Callable[[], None]):
        self._listeners.append(fn)

    def _emit(self):
        for fn in self._listeners:
            fn()

    # ------------------------------------------------------------- undo
    def checkpoint(self):
        """Call before a user-visible change. One checkpoint = one undo step."""
        if getattr(self, "_suspend_cp", 0):           # inside a copilot turn: the caller made one already
            return
        self._undo.append(json.dumps(self.to_dict()))
        self._undo = self._undo[-200:]
        self._redo.clear()

    def undo(self) -> bool:
        if not self._undo:
            return False
        self._redo.append(json.dumps(self.to_dict()))
        self._load_dict(json.loads(self._undo.pop()))
        self.rebuild()
        return True

    def redo(self) -> bool:
        if not self._redo:
            return False
        self._undo.append(json.dumps(self.to_dict()))
        self._load_dict(json.loads(self._redo.pop()))
        self.rebuild()
        return True

    # ------------------------------------------------------------- layers
    def layer(self, name: str) -> Layer:
        for l in self.layers:
            if l.name == name:
                return l
        raise FeatureError(f"No layer named '{name}'")

    def has_layer(self, name: str) -> bool:
        return any(l.name == name for l in self.layers)

    def _check_unlocked(self, f: Feature):
        if self.has_layer(f.layer) and self.layer(f.layer).locked:
            raise FeatureError(f"{f.name} is on locked layer '{f.layer}' — unlock the layer to change it")

    def add_layer(self, name: str, color: str | None = None, checkpoint=True, activate=False) -> Layer:
        name = name.strip()
        if not valid_name(name):
            raise FeatureError("Layer names: 1–41 characters, no slashes")
        if self.has_layer(name):
            raise FeatureError(f"Layer '{name}' already exists")
        if checkpoint:
            self.checkpoint()
        l = Layer(name, color or PALETTE[len(self.layers) % len(PALETTE)])
        self.layers.append(l)
        if activate:
            self.active_layer = name
        self.dirty = True
        self._emit()
        return l

    def set_layer(self, name: str, checkpoint=True, **props):
        """props: color, visible, locked, in_model, new_name, active."""
        l = self.layer(name)
        if checkpoint:
            self.checkpoint()
        rebuild = False
        if "new_name" in props and props["new_name"] and props["new_name"] != name:
            new = props["new_name"].strip()
            if not valid_name(new) or self.has_layer(new):
                raise FeatureError(f"Can't rename layer to '{new}'")
            for f in self.features:
                if f.layer == name:
                    f.layer = new
            if self.active_layer == name:
                self.active_layer = new
            l.name = new
        for k in ("color", "visible", "locked"):
            if k in props and props[k] is not None:
                setattr(l, k, props[k])
        if "in_model" in props and props["in_model"] is not None and props["in_model"] != l.in_model:
            l.in_model = bool(props["in_model"])
            rebuild = True
        if props.get("active"):
            self.active_layer = l.name
        self.dirty = True
        if rebuild:
            self.rebuild()
        else:
            self._emit()

    def delete_layer(self, name: str, move_to: str = DEFAULT, checkpoint=True):
        if name == DEFAULT:
            raise FeatureError("The Default layer can't be deleted")
        l = self.layer(name)
        target = self.layer(move_to)
        if l.locked:
            raise FeatureError(f"Layer '{name}' is locked")
        if checkpoint:
            self.checkpoint()
        for f in self.features:
            if f.layer == name:
                f.layer = target.name
        self.layers.remove(l)
        if self.active_layer == name:
            self.active_layer = target.name
        self.dirty = True
        self.rebuild()

    def move_to_layer(self, names: list[str], layer: str, checkpoint=True):
        target = self.layer(layer)
        if target.locked:
            raise FeatureError(f"Layer '{layer}' is locked")
        feats = [self.get(n) for n in names]
        for f in feats:
            self._check_unlocked(f)
        if checkpoint:
            self.checkpoint()
        for f in feats:
            f.layer = target.name
        self.dirty = True
        self.rebuild()

    def layer_counts(self) -> dict[str, int]:
        out = {l.name: 0 for l in self.layers}
        for f in self.features:
            out[f.layer] = out.get(f.layer, 0) + 1
        return out

    # ------------------------------------------------------------ objects
    def obj(self, name: str) -> Obj:
        o = self._oidx().get(name)
        if o is not None and o.name == name:
            return o
        for o in self.objects:
            if o.name == name:
                self._names_changed()
                return o
        raise FeatureError(f"No object named '{name}'")

    def has_obj(self, name: str) -> bool:
        o = self._oidx().get(name)
        if o is not None and o.name == name:
            return True
        return False

    # name indexes: big imports add thousands of objects/features — keep lookups O(1)
    def _oidx(self) -> dict:
        c = self.__dict__.get("_oidx_c")
        if c is None or c[0] is not self.objects or c[1] != len(self.objects):
            c = [self.objects, len(self.objects), {o.name: o for o in self.objects}]
            self.__dict__["_oidx_c"] = c
        return c[2]

    def _fnames(self) -> set:
        c = self.__dict__.get("_fidx_c")
        if c is None or c[0] is not self.features or c[1] != len(self.features):
            c = [self.features, len(self.features), {f.name for f in self.features}]
            self.__dict__["_fidx_c"] = c
        return c[2]

    def _names_changed(self):
        self.__dict__["_oidx_c"] = None
        self.__dict__["_fidx_c"] = None

    def _noted_add(self, kind, item):
        """Keep an index valid after appending one item (instead of rebuilding it)."""
        key, lst = ("_oidx_c", self.objects) if kind == "obj" else ("_fidx_c", self.features)
        c = self.__dict__.get(key)
        if c is not None and c[0] is lst and c[1] == len(lst) - 1:
            c[1] += 1
            if kind == "obj":
                c[2][item.name] = item
            else:
                c[2].add(item.name)

    def add_object(self, name: str | None = None, color: str | None = None, checkpoint=True,
                   activate=True, pose=None) -> Obj:
        if checkpoint:
            self.checkpoint()
        if not name:
            n = 1
            while self.has_obj(f"Object{n}"):
                n += 1
            name = f"Object{n}"
        name = name.strip()
        if not valid_name(name) or self.has_obj(name):
            raise FeatureError(f"Object name '{name}' is invalid or taken")
        o = Obj(name, color or PALETTE[(len(self.objects) + 1) % len(PALETTE)],
                pose=list(pose) if pose else [0.0] * 6)
        self.objects.append(o)
        self._noted_add("obj", o)
        if activate:
            self.active_object = name
        self.dirty = True
        self._emit()
        return o

    def set_object(self, name: str, checkpoint=True, **props):
        """props: color, visible, fixed, new_name, active, pose."""
        o = self.obj(name)
        if checkpoint:
            self.checkpoint()
        new = props.get("new_name")
        if new and new != name:
            new = new.strip()
            if not valid_name(new) or self.has_obj(new):
                raise FeatureError(f"Can't rename object to '{new}'")
            for f in self.features:
                if f.obj == name:
                    f.obj = new
            for m in self.mates:
                for side in (m.a, m.b):
                    if side["obj"] == name:
                        side["obj"] = new
                for v in m.extra.values():
                    if isinstance(v, dict) and v.get("obj") == name:
                        v["obj"] = new
            for pat in self.comp_patterns:
                pat["seeds"] = [new if x == name else x for x in pat["seeds"]]
                pat["instances"] = [new if x == name else x for x in pat.get("instances", [])]
            for x in self.objects:
                if x.instance_of == name:
                    x.instance_of = new
            if self.active_object == name:
                self.active_object = new
            o.name = new
            self._names_changed()
        for k in ("color", "visible", "fixed"):
            if props.get(k) is not None:
                setattr(o, k, props[k])
        if props.get("pose") is not None:
            o.pose = [float(v) for v in props["pose"]]
        if props.get("active"):
            self.active_object = o.name
        self.dirty = True
        if props.get("pose") is not None or props.get("fixed") is not None:
            self.solve_assembly()
        self._emit()

    def place_object(self, name: str, translate=(0, 0, 0), axis=(0, 0, 1), angle_deg=0.0,
                     about=(0, 0, 0), checkpoint=True):
        """Move an object rigidly in world space (translate, then rotate about a world axis
        through `about`). Mates then pull connected objects along."""
        import numpy as np
        o = self.obj(name)
        if checkpoint:
            self.checkpoint()
        ax = np.asarray(axis, float)
        ax = ax / (np.linalg.norm(ax) or 1)
        rv = ax * np.radians(angle_deg)
        about = np.asarray(about, float)
        move = list(np.asarray(translate, float)) + [0, 0, 0]
        if angle_deg:
            R = asm.rotmat(rv)
            t = about - R @ about
            rot = list(t) + list(rv)
            o.pose = asm.compose(rot, asm.compose(move, o.pose))
        else:
            o.pose = asm.compose(move, o.pose)
        self.dirty = True
        if not o.fixed:
            self.solve_assembly(keep=name)
        self._emit()

    # ------------------------------------------------ object manipulation
    FLOOR_Z = 0.0

    def object_local_points(self, name: str):
        """Sample points (vertices + edge polylines) of an object's body in its own coordinates."""
        import numpy as np

        from . import snap3d
        b = self.result.bodies.get(name)
        if b is None:
            raise FeatureError(f"{name} has no geometry yet")
        c = snap3d.candidates(b, key=("pts", name, id(b)))
        parts = [c["pts"]] + [np.asarray(e, float) for e in c["edges"]]
        return np.vstack([q for q in parts if len(q)])

    def object_world_points(self, name: str, pose=None):
        import numpy as np
        P = self.object_local_points(name)
        M = asm.pose_matrix(self.obj(name).pose if pose is None else pose)
        return P @ M[:3, :3].T + M[:3, 3]

    def object_center(self, name: str):
        """World bounding-box centre of an object."""
        W = self.object_world_points(name)
        return (W.min(0) + W.max(0)) / 2

    def _movable(self, name):
        """Explicit placement commands work on fixed (grounded) objects too: the ground moves and the
        mated objects follow."""
        o = self.obj(name)
        if self.result.bodies.get(name) is None:
            raise FeatureError(f"{name} has no geometry yet")
        return o

    @staticmethod
    def rotated_pose(pose, R, pivot):
        """Pose after rotating by world matrix R about world point `pivot`."""
        import numpy as np
        pivot = np.asarray(pivot, float)
        rot = list(pivot - R @ pivot) + list(asm.rotvec(R))
        return asm.compose(rot, pose)

    def _settle(self, name):
        self.dirty = True
        self.solve_assembly(keep=name)
        self._emit()

    def rotate_object(self, name: str, axis=(0, 0, 1), angle_deg=90.0, pivot="center", on_floor=False,
                      checkpoint=True):
        """Rotate an object about a world axis through its centre ("center"), its origin ("origin")
        or a world point. on_floor: keep its lowest point where it was (rotate about the floor)."""
        import numpy as np
        o = self._movable(name)
        if checkpoint:
            self.checkpoint()
        ax = np.asarray(axis, float)
        ax = ax / (np.linalg.norm(ax) or 1.0)
        R = asm.rotmat(ax * np.radians(angle_deg))
        if isinstance(pivot, str):
            pv = self.object_center(name) if pivot == "center" else np.asarray(o.pose[:3], float)
        else:
            pv = np.asarray(pivot, float)
        zmin0 = self.object_world_points(name)[:, 2].min() if on_floor else None
        o.pose = self.rotated_pose(o.pose, R, pv)
        if on_floor:
            dz = zmin0 - self.object_world_points(name)[:, 2].min()
            o.pose = [o.pose[0], o.pose[1], o.pose[2] + float(dz)] + list(o.pose[3:])
        self._settle(name)

    UP_AXES = {"+Z": None, "-Z": ((1, 0, 0), 180.0), "+Y": ((1, 0, 0), 90.0), "-Y": ((1, 0, 0), -90.0),
               "+X": ((0, 1, 0), -90.0), "-X": ((0, 1, 0), 90.0)}

    def up_poses(self, names, up="+Z", floor=True, base=None) -> dict:
        """New poses that turn a group of objects so that `up` (the axis that was up in the source CAD) points
        along +Z, rotating the group about its common centre (parts keep their relative placement); floor:
        then set the group's lowest point on z = 0. base: starting poses {name: pose} (default: current)."""
        import numpy as np
        if up not in self.UP_AXES:
            raise FeatureError(f"up must be one of {list(self.UP_AXES)}")
        base = {n: list(base[n]) if base and n in base else list(self.obj(n).pose) for n in names}
        pts = {n: self.object_world_points(n, base[n]) for n in names}
        allp = np.vstack([p for p in pts.values() if len(p)]) if pts else np.zeros((1, 3))
        center = (allp.min(0) + allp.max(0)) / 2
        rot = self.UP_AXES[up]
        R = asm.rotmat(np.asarray(rot[0], float) * np.radians(rot[1])) if rot else np.eye(3)
        out = {n: (self.rotated_pose(base[n], R, center) if rot else list(base[n])) for n in names}
        if floor:
            zmin = min(self.object_world_points(n, out[n])[:, 2].min() for n in names)
            dz = self.FLOOR_Z - float(zmin)
            out = {n: [p[0], p[1], p[2] + dz] + list(p[3:]) for n, p in out.items()}
        return out

    def set_up_axis(self, names, up="+Z", floor=True, checkpoint=True):
        """Apply up_poses (one undo step)."""
        poses = self.up_poses(names, up, floor)
        if checkpoint:
            self.checkpoint()
        for n, p in poses.items():
            self.obj(n).pose = [float(v) for v in p]
        self.dirty = True
        self._emit()
        return poses

    def drop_to_floor(self, name: str, floor: float | None = None, checkpoint=True):
        """Translate an object straight down/up so its lowest point sits on the floor (z = 0)."""
        o = self._movable(name)
        if checkpoint:
            self.checkpoint()
        floor = self.FLOOR_Z if floor is None else floor
        dz = floor - self.object_world_points(name)[:, 2].min()
        o.pose = [o.pose[0], o.pose[1], o.pose[2] + float(dz)] + list(o.pose[3:])
        self._settle(name)

    def face_to_floor(self, name: str, face_index: int, local_point=None, floor: float | None = None,
                      checkpoint=True):
        """Turn an object so the given face points straight down, then set it on the floor."""
        import numpy as np
        from build123d import Vector
        o = self._movable(name)
        b = self.result.bodies.get(name)
        faces = list(b.faces()) if b is not None else []
        if not 0 <= int(face_index) < len(faces):
            raise FeatureError(f"{name} has no face {face_index}")
        f = faces[int(face_index)]
        try:
            nv = f.normal_at(Vector(*local_point)) if local_point is not None else f.normal_at()
        except Exception:  # noqa: BLE001
            nv = f.normal_at()
        R0 = asm.rotmat(o.pose[3:6])
        n = R0 @ np.array([nv.X, nv.Y, nv.Z])
        n = n / (np.linalg.norm(n) or 1.0)
        down = np.array([0.0, 0.0, -1.0])
        if checkpoint:
            self.checkpoint()
        c = float(np.clip(n @ down, -1.0, 1.0))
        if c < 1 - 1e-9:
            axis = np.cross(n, down)
            if np.linalg.norm(axis) < 1e-9:
                axis = np.array([1.0, 0.0, 0.0])
            axis = axis / np.linalg.norm(axis)
            R = asm.rotmat(axis * np.arccos(c))
            o.pose = self.rotated_pose(o.pose, R, self.object_center(name))
        floor = self.FLOOR_Z if floor is None else floor
        dz = floor - self.object_world_points(name)[:, 2].min()
        o.pose = [o.pose[0], o.pose[1], o.pose[2] + float(dz)] + list(o.pose[3:])
        self._settle(name)

    def set_orientation(self, name: str, rx=0.0, ry=0.0, rz=0.0, keep_center=True, checkpoint=True):
        """Absolute orientation from X→Y→Z angles (degrees, world axes); keeps the centre in place."""
        import numpy as np
        o = self._movable(name)
        if checkpoint:
            self.checkpoint()
        c0 = self.object_center(name) if keep_center else None
        R = (asm.rotmat([0, 0, np.radians(rz)]) @ asm.rotmat([0, np.radians(ry), 0])
             @ asm.rotmat([np.radians(rx), 0, 0]))
        o.pose = list(o.pose[:3]) + list(asm.rotvec(R))
        if keep_center:
            d = c0 - self.object_center(name)
            o.pose = [o.pose[i] + float(d[i]) for i in range(3)] + list(o.pose[3:])
        self._settle(name)

    @staticmethod
    def orientation_angles(pose):
        """X→Y→Z angles (degrees) of a pose's rotation (inverse of set_orientation)."""
        import math

        import numpy as np
        R = asm.rotmat(pose[3:6])
        ry = math.asin(-max(-1.0, min(1.0, R[2, 0])))
        if abs(math.cos(ry)) > 1e-6:
            rx = math.atan2(R[2, 1], R[2, 2])
            rz = math.atan2(R[1, 0], R[0, 0])
        else:
            rx = math.atan2(-R[1, 2], R[1, 1])
            rz = 0.0
        return [float(np.degrees(v)) for v in (rx, ry, rz)]

    def delete_object(self, name: str, checkpoint=True):
        if len(self.objects) == 1:
            raise FeatureError("A document needs at least one object")
        o = self.obj(name)
        feats = [f for f in self.features if f.obj == name]
        for f in feats:
            self._check_unlocked(f)
        if checkpoint:
            self.checkpoint()
        self.features = [f for f in self.features if f.obj != name]
        for pat in list(self.comp_patterns):                 # its pattern instances go with it
            if name in pat["seeds"]:
                pat["seeds"].remove(name)
                self._drop_instances(pat)
                if pat["seeds"]:
                    self._regen_pattern(pat)
                else:
                    self.comp_patterns.remove(pat)
            elif name in pat.get("instances", []):
                pat["instances"].remove(name)
        self.objects = [x for x in self.objects if x.instance_of != name]
        self._names_changed()
        self.mates = [m for m in self.mates if name not in m.objects()]
        names = {m.name for m in self.mates}
        self.physics.drives = [d for d in self.physics.drives if d.mate in names]
        self.objects.remove(o)
        if self.active_object == name:
            self.active_object = self.objects[0].name
        self.dirty = True
        self.rebuild()

    # -------------------------------------------------------------- mates
    def mate(self, name: str) -> Mate:
        for m in self.mates:
            if m.name == name:
                return m
        raise FeatureError(f"No mate named '{name}'")

    def ref_on(self, obj: str, spec) -> dict:
        """A mate reference on an object (in its own coordinates). spec forms:
        stored ref dict · face index (int) · {"face_index": i} · {"face": "top"|"bottom"|…} (extreme face)
        · {"face_near": [x,y,z], "normal": [nx,ny,nz]?} · {"axis": "<feature>", "near": [x,y,z]?, "radius": r?}
        · {"axis_near": [x,y,z]} (closest cylinder)
        · {"point": [x,y,z]}."""
        from .selectors import select_face
        body = self.result.bodies.get(obj)
        if isinstance(spec, dict) and "kind" in spec:
            return dict(spec)
        special = self._ref_special(obj, spec)
        if special is not None:
            return special
        if body is None:
            raise FeatureError(f"Object '{obj}' has no body yet")
        faces = list(body.faces())
        if isinstance(spec, int):
            spec = {"face_index": spec}
        if "face_index" in spec:
            i = int(spec["face_index"])
            if not 0 <= i < len(faces):
                raise FeatureError(f"{obj} has no face {i}")
            return asm.ref_from_face(faces[i], i, spec.get("pick"))
        if "face_near" in spec:
            import numpy as np
            p = np.asarray(spec["face_near"], float)
            want = np.asarray(spec["normal"], float) if spec.get("normal") else None
            best, bd = None, float("inf")
            for i, g in enumerate(faces):
                if want is not None:
                    if g.geom_type.name != "PLANE":
                        continue
                    nn = g.normal_at()
                    if float(np.dot([nn.X, nn.Y, nn.Z], want / np.linalg.norm(want))) < 0.99:
                        continue
                dist = g.distance_to(tuple(p))
                if dist < bd:
                    best, bd = i, dist
            if best is None:
                raise FeatureError(f"No matching face near {list(p)} on {obj}")
            return asm.ref_from_face(faces[best], best, tuple(p) if want is not None else None)
        if "face" in spec:
            f = select_face(body, spec["face"])
            i = next((k for k, g in enumerate(faces) if g.is_same(f)), None)
            return asm.ref_from_face(f, i)
        if "axis" in spec or "axis_near" in spec:
            import numpy as np
            owner = spec.get("axis")
            near = spec.get("near", spec.get("axis_near"))
            fl = self.result.face_owner_by_obj.get(obj, [])
            cands = [(i, g) for i, g in enumerate(faces) if g.geom_type.name in ("CYLINDER", "CONE")]
            if owner not in (None, "", "any"):
                owned = [(i, g) for i, g in cands if (fl[i] if i < len(fl) else None) == owner]
                if not owned:
                    raise FeatureError(f"No cylindrical face from '{owner}' on {obj}")
                cands = owned
            if not cands:
                raise FeatureError(f"{obj} has no cylindrical faces")
            if near is not None:
                p = tuple(float(v) for v in near)
                cands.sort(key=lambda ig: ig[1].distance_to(p))
            elif "radius" in spec:
                cands.sort(key=lambda ig: abs(asm.ref_from_face(ig[1]).get("r", 0) - float(spec["radius"])))
            i, g = cands[0]
            return asm.ref_from_face(g, i)
        if "point" in spec:
            return {"kind": "point", "p": [float(v) for v in spec["point"]], "face": None}
        raise FeatureError(f"Can't understand reference {spec}")

    def add_mate(self, mtype: str, a: dict, b: dict, value: float = 0.0, flip: bool | None = None,
                 name: str | None = None, checkpoint=True, value2: float | None = None,
                 extra: dict | None = None) -> Mate:
        """a/b: {"obj": name, "ref": spec} — see ref_on for spec forms. value2: maximum of a limit mate.
        extra: width {"a2", "b2"} / symmetric {"c"} as {"obj", "ref"}; path {"percent"}; slot {"along":
        free|center|percent, "percent"}."""
        if a["obj"] == b["obj"]:
            raise FeatureError("A mate joins two different objects")
        ra, rb = self.ref_on(a["obj"], a["ref"]), self.ref_on(b["obj"], b["ref"])
        try:
            asm.check_mate_kinds(mtype, ra["kind"], rb["kind"])
        except AssemblyError as e:
            raise FeatureError(str(e)) from e
        ext = {}
        for k, v in (extra or {}).items():
            if k in asm.EXTRA_REFS.get(mtype, ()):
                r = self.ref_on(v["obj"], v["ref"])
                if r["kind"] != "plane" and (mtype == "width" or k == "c"):
                    raise FeatureError(f"{mtype}: '{k}' must be a flat face")
                ext[k] = {"obj": v["obj"], "ref": r}
            else:
                ext[k] = v
        missing = [k for k in asm.EXTRA_REFS.get(mtype, ()) if k not in ext]
        if missing:
            raise FeatureError(f"A {mtype} mate also needs {', '.join(missing)}")
        if mtype in asm.RANGE_MATES:
            lo = float(value)
            hi = float(value2 if value2 is not None else value)
            if hi < lo:
                raise FeatureError("Limit mate: maximum is below the minimum")
        if checkpoint:
            self.checkpoint()
        n = 1
        while any(m.name == f"{mtype.title()}{n}" for m in self.mates):
            n += 1
        m = Mate(name or f"{mtype.title().replace('_', '')}{n}", mtype, {"obj": a["obj"], "ref": ra},
                 {"obj": b["obj"], "ref": rb}, float(value), False, False,
                 float(value2 if value2 is not None else value), ext)
        if flip is None:
            # choose the orientation closest to how the parts sit now
            import numpy as np
            A = asm.world_ref(ra, self.obj(a["obj"]).pose)
            B = asm.world_ref(rb, self.obj(b["obj"]).pose)
            if ra["kind"] == rb["kind"] == "plane":
                m.flip = bool(np.dot(A["n"], B["n"]) > 0)
        else:
            m.flip = bool(flip)
        if mtype == "tangent" and flip is None and {ra["kind"], rb["kind"]} == {"plane", "axis"}:
            import numpy as np
            A = asm.world_ref(ra, self.obj(a["obj"]).pose)
            B = asm.world_ref(rb, self.obj(b["obj"]).pose)
            P, X = (A, B) if ra["kind"] == "plane" else (B, A)
            m.flip = bool(np.dot(P["n"], X["p"] - P["p"]) < 0)
        if mtype in ("gear", "rack", "screw"):
            pa_, pb_ = self.obj(a["obj"]).pose, self.obj(b["obj"]).pose
            m.a["R0"] = asm.rotmat(pa_[3:6]).tolist()
            m.b["R0"] = asm.rotmat(pb_[3:6]).tolist()
            m.a["t0"] = list(pa_[:3])
            m.b["t0"] = list(pb_[:3])
        if mtype == "lock":
            import numpy as np
            Ma = asm.pose_matrix(self.obj(a["obj"]).pose)
            Mb = asm.pose_matrix(self.obj(b["obj"]).pose)
            m.a["rel"] = (np.linalg.inv(Ma) @ Mb).tolist()
        self.mates.append(m)
        before = {o.name: list(o.pose) for o in self.objects}
        res = self.solve_assembly()
        if res is not None and m.name in res.failed:
            self.mates.remove(m)
            for o in self.objects:
                o.pose = before[o.name]
            self.solve_assembly()
            if checkpoint:
                self._undo.pop()
            raise FeatureError(f"{m.name} conflicts with the existing mates (over-defined) — not added")
        self.dirty = True
        self._emit()
        return m

    def set_mate(self, name: str, checkpoint=True, **props):
        m = self.mate(name)
        if checkpoint:
            self.checkpoint()
        for k in ("value", "flip", "suppressed", "value2"):
            if props.get(k) is not None:
                setattr(m, k, type(getattr(m, k))(props[k]))
        self.dirty = True
        self.solve_assembly()
        self._emit()

    def delete_mate(self, name: str, checkpoint=True):
        m = self.mate(name)
        if checkpoint:
            self.checkpoint()
        self.mates.remove(m)
        self.physics.drives = [d for d in self.physics.drives if d.mate != name]
        self.dirty = True
        self.solve_assembly()
        self._emit()

    def _mate_refs(self) -> dict:
        refs = {}
        for m in self.mates:
            sides = [(m.a, "a"), (m.b, "b")] + [(m.extra[k], k) for k in asm.EXTRA_REFS.get(m.type, ())
                                                if k in m.extra]
            for side, key in sides:
                if side["ref"].get("kind") in ("path", "profile", "slot"):
                    refs[(m.name, key)] = side["ref"]
                    continue
                body = self.result.bodies.get(side["obj"])
                r = asm.resolve_ref(body, side["ref"]) if body is not None else side["ref"]
                side["ref"] = {**r}
                refs[(m.name, key)] = r
        return refs

    def solve_assembly(self, drag=None, keep: str | None = None, apply=True):
        """Solve mates; update object poses. keep: treat this object as fixed for this solve
        (it was just placed by hand). drag: (obj, local point, world target)."""
        if not self.mates and not (drag and self._rigid_groups()):
            if self.result is not None:
                self.result.assembly = None
            self._update_pattern_poses()
            if drag and apply:                       # no mates: the dragged object just translates
                name, lp, tgt = drag
                o = self.obj(name)
                if not o.fixed:
                    M = asm.pose_matrix(o.pose)
                    cur = (M @ np.append(np.asarray(lp, float), 1.0))[:3]
                    if isinstance(tgt, dict) and "ray" in tgt:
                        ro, rd = (np.asarray(v, float) for v in tgt["ray"])
                        rd = rd / np.linalg.norm(rd)
                        tgt = ro + rd * float((cur - ro) @ rd)
                    delta = np.asarray(tgt, float) - cur
                    o.pose = [float(o.pose[i] + delta[i]) for i in range(3)] + list(o.pose[3:])
            return None
        objs = [Obj(o.name, o.color, o.visible, o.fixed or o.name == keep or bool(o.pattern), list(o.pose))
                for o in self.objects]
        gone = {o.name for o in self.objects if o.suppressed}
        mates = [m for m in self.mates if not set(m.objects()) & gone]
        refs = self._mate_refs()
        implicit = []
        for grp in self._rigid_groups():                     # rigid subassemblies move as one body
            anchor = grp[0]
            Ma = asm.pose_matrix(self.obj(anchor).pose)
            for other in grp[1:]:
                Mb = asm.pose_matrix(self.obj(other).pose)
                lm = Mate(f"_rigid_{anchor}_{other}", "lock", {"obj": anchor, "ref": {"kind": "point", "p": [0, 0, 0]},
                                                                "rel": (np.linalg.inv(Ma) @ Mb).tolist()},
                          {"obj": other, "ref": {"kind": "point", "p": [0, 0, 0]}})
                implicit.append(lm)
                refs[(lm.name, "a")] = lm.a["ref"]
                refs[(lm.name, "b")] = lm.b["ref"]
        if not mates and not implicit:
            if self.result is not None:
                self.result.assembly = None
            self._update_pattern_poses()
            return None
        if not mates and not drag:
            self._update_pattern_poses()
            if self.result is not None:
                self.result.assembly = None
            return None
        try:
            res = asm.solve(objs, mates + implicit, refs, drag=drag)
        except (AssemblyError, KeyError) as e:
            self.result.errors["Assembly"] = str(e)
            return None
        for lm in implicit:
            res.residuals.pop(lm.name, None)
        res.failed = [f for f in res.failed if not f.startswith("_rigid_")]
        res.ok = not res.failed
        if apply:
            for o in self.objects:
                o.pose = res.poses[o.name]
            self._update_pattern_poses()
            self.result.assembly = res
        return res

    def drive(self, mate_name: str, start: float, end: float, steps: int = 24) -> list[dict]:
        """Motion study: sweep a mate's value, solving each step. Returns [{obj: pose}] frames;
        the document is left at the final step."""
        m = self.mate(mate_name)
        if not asm.MATE_TYPES[m.type][1]:
            raise FeatureError(f"{m.name} has no value to drive (use a distance or angle mate)")
        frames = []
        for i in range(steps + 1):
            m.value = start + (end - start) * i / max(steps, 1)
            res = self.solve_assembly()
            frames.append({o.name: list(o.pose) for o in self.objects} | {"_value": m.value,
                                                                           "_ok": bool(res and res.ok)})
        self.dirty = True
        self._emit()
        return frames

    def placed_bodies(self) -> dict:
        """{object: body moved to its pose} for visible-or-not objects that have bodies."""
        out = {}
        for o in self.objects:
            b = self.result.bodies.get(o.name)
            if b is not None:
                out[o.name] = b.moved(asm.location_of(o.pose)) if any(o.pose) else b
        return out

    def interference(self, min_volume=1e-3) -> list[tuple[str, str, float]]:
        bodies = list(self.placed_bodies().items())
        hits = []
        for i in range(len(bodies)):
            for j in range(i + 1, len(bodies)):
                (na, a), (nb, b) = bodies[i], bodies[j]
                try:
                    ba, bb = a.bounding_box(), b.bounding_box()
                    if (ba.max.X < bb.min.X or bb.max.X < ba.min.X or ba.max.Y < bb.min.Y or
                            bb.max.Y < ba.min.Y or ba.max.Z < bb.min.Z or bb.max.Z < ba.min.Z):
                        continue
                    v = (a & b).volume
                except Exception:  # noqa: BLE001
                    continue
                if v > min_volume:
                    hits.append((na, nb, round(v, 3)))
        return hits

    # ------------------------------------------------------------ physics
    def set_physics(self, checkpoint=True, **props):
        """props: gravity [x,y,z] m/s², duration s, timestep s, friction, collisions, fps."""
        from .physics import COLLISIONS
        vals = {}
        for k in ("gravity", "duration", "timestep", "friction", "collisions", "fps"):
            if props.get(k) is None:
                continue
            v = props[k]
            if k == "gravity":
                v = [float(x) for x in v]
                if len(v) != 3:
                    raise FeatureError("gravity is [x, y, z] in m/s²")
            elif k == "collisions":
                if v not in COLLISIONS:
                    raise FeatureError(f"collisions: one of {COLLISIONS}")
            elif k == "fps":
                v = max(5, min(240, int(v)))
            else:
                v = float(v)
                if k in ("duration", "timestep") and v <= 0:
                    raise FeatureError(f"{k} must be > 0")
                if k == "timestep" and v > 0.01:
                    raise FeatureError("timestep must be ≤ 0.01 s")
            vals[k] = v
        if checkpoint:
            self.checkpoint()
        for k, v in vals.items():
            setattr(self.physics, k, v)
        self.dirty = True
        self._emit()

    def set_drive(self, mate: str, kind: str | None, value: float = 0.0, damping: float = 0.0,
                  max_effort: float = 0.0, checkpoint=True):
        """Motor/spring on a mate. kind None removes it. One drive per mate."""
        self.mate(mate)
        if kind is not None and kind not in DRIVE_KINDS:
            raise FeatureError(f"Drive kind must be one of {DRIVE_KINDS}")
        if checkpoint:
            self.checkpoint()
        self.physics.drives = [d for d in self.physics.drives if d.mate != mate]
        if kind is not None:
            self.physics.drives.append(Drive(mate, kind, float(value), float(damping), float(max_effort)))
        self.dirty = True
        self._emit()

    def simulate(self, progress=None, cancel=None):
        from .physics import simulate
        try:
            return simulate(self, self.physics, progress=progress, cancel=cancel)
        except PhysicsError as e:
            raise FeatureError(str(e)) from e

    def apply_frame(self, frame: dict, checkpoint=True):
        """Keep a simulated state: object poses from a frame."""
        if checkpoint:
            self.checkpoint()
        for o in self.objects:
            if o.name in frame and not o.fixed:
                o.pose = [float(v) for v in frame[o.name]]
        # driven distance/angle mates take the values they reached, so the solver keeps this state
        import math
        refs = self._mate_refs()
        P = {o.name: o.pose for o in self.objects}
        for m in self.mates:
            if m.suppressed or m.type not in ("distance", "angle"):
                continue
            try:
                r = asm._eval(m, refs, P)
            except Exception:  # noqa: BLE001
                continue
            if m.type == "distance":
                m.value = float(m.value + r[-1])
            else:
                m.value = math.degrees(math.acos(max(-1.0, min(1.0, r[0] + math.cos(math.radians(m.value))))))
        self.dirty = True
        self.rebuild()

    # ------------------------------------------------------------ studies
    STUDY_ITEM_PARAMS = {
        "fixtures": {"fixed": (), "roller": (), "pin": ("slide",), "elastic": ("stiffness", "normal_only")},
        "loads": {"force": ("value", "direction"), "pressure": ("value",), "torque": ("value", "direction"),
                  "bearing": ("value", "direction"), "gravity": ("value", "direction"),
                  "centrifugal": ("value", "point", "direction")},
        "thermal": {"temperature": ("value",), "convection": ("h", "ambient"), "heat_flux": ("value",),
                    "heat_power": ("value",), "radiation": ("emissivity", "ambient"), "heat_volume": ("value",)},
    }

    def add_bolt(self, study: str, a_object: str, a_face, b_object: str, b_face, preload: float = 0.0,
                 diameter: float = 0.0, grip: float = 0.0, grade: str = "8.8", checkpoint=True) -> dict:
        """Bolt connector in an assembly study between a hole face on one object and one on another."""
        from ..sim.studies import face_refs
        st = self.study(study)
        if st["type"] != "assembly":
            raise FeatureError("Bolts go in an assembly study")
        try:
            ra = face_refs(self, a_object, [a_face])
            rb = face_refs(self, b_object, [b_face])
        except Exception as e:  # noqa: BLE001
            raise FeatureError(str(e)) from e
        if ra[0].get("kind") != "axis" or rb[0].get("kind") != "axis":
            raise FeatureError("Pick the hole (cylindrical) faces for a bolt")
        if checkpoint:
            self.checkpoint()
        st.setdefault("connectors", [])
        item = {"type": "bolt", "name": f"Bolt {len(st['connectors']) + 1}", "a_object": a_object, "a_refs": ra,
                "b_object": b_object, "b_refs": rb, "preload": float(preload), "diameter": float(diameter),
                "grip": float(grip), "grade": grade}
        st["connectors"].append(item)
        self.dirty = True
        self._emit()
        return item

    def design_study(self, feature: str, param: str, values, study: str | None = None, metrics=None,
                     progress=None) -> dict:
        """Parameter sweep: each value is applied, rebuilt and measured (mass + the study's results)."""
        from ..sim.advanced import design_study
        from ..sim.fea import SimError
        if study:
            self.study(study)
        if not values:
            raise FeatureError("Give the values to try")
        try:
            out = design_study(self, feature, param, list(values), study, metrics, progress)
        except SimError as e:
            raise FeatureError(str(e)) from e
        except (KeyError, ValueError) as e:
            raise FeatureError(str(e)) from e
        self.last_design_study = out
        return out

    def topology_to_body(self, study: str, name: str | None = None) -> str:
        """Turn a topology result into a new object (meshed solid) for redesign or export."""
        from ..io.formats import mesh_to_solid, shape_to_text, store_blob
        from ..sim.advanced import topology_surface
        res = self.sim_results.get(study)
        if res is None or res.kind != "topology":
            raise FeatureError(f"Run topology study '{study}' first")
        V, F = topology_surface(res)
        try:
            shape = mesh_to_solid(V, F)
        except Exception as e:  # noqa: BLE001
            raise FeatureError(f"Couldn't make a solid: {e}") from e
        st = self.study(study)
        self.checkpoint()
        o = self.add_object(self._unique_obj_name(name or f"{st['object']} optimised"), "#6fbf73", checkpoint=False,
                            activate=False, pose=list(self.obj(st["object"]).pose))
        o.material = self.material_of(st["object"]).name
        self.add("import", {"data": store_blob(shape_to_text(shape)), "source": f"topology {study}",
                            "format": "mesh"}, name=self._unique_feature_name(o.name.replace(" ", "_")),
                 rebuild=True, checkpoint=False, obj=o.name)
        return o.name

    def study(self, name: str) -> dict:
        for st in self.studies:
            if st["name"] == name:
                return st
        raise FeatureError(f"No study named '{name}'")

    def add_study(self, stype: str, name: str | None = None, obj: str | None = None, checkpoint=True,
                  **options) -> dict:
        from ..sim import studies as S
        try:
            st = S.new_study(self, stype, name, obj, **options)
        except S.SimError as e:
            raise FeatureError(str(e)) from e
        if checkpoint:
            self.checkpoint()
        self.studies.append(st)
        self.dirty = True
        self._emit()
        return st

    def set_study(self, name: str, checkpoint=True, **props):
        """props: options of the study type, mesh_size, object, new_name."""
        from ..sim.studies import OPTIONS
        st = self.study(name)
        opts = OPTIONS.get(st["type"], {})
        for k in props:
            if k not in opts and k not in ("mesh_size", "object", "new_name"):
                raise FeatureError(f"{st['type']} study has no option '{k}'. Options: {list(opts)}")
        if props.get("object"):
            self.obj(props["object"])
        if checkpoint:
            self.checkpoint()
        for k, v in props.items():
            if k == "mesh_size":
                st["mesh_size"] = float(v) if v else None
            elif k == "object":
                st["object"] = v
            elif k == "new_name":
                if v and v != name:
                    if any(x["name"] == v for x in self.studies):
                        raise FeatureError(f"Study name '{v}' already used")
                    st["name"] = v
                    if name in self.sim_results:
                        self.sim_results[v] = self.sim_results.pop(name)
            else:
                st["options"][k] = _coerce_opt(opts[k][0], v)
        self.dirty = True
        self._emit()

    def delete_study(self, name: str, checkpoint=True):
        st = self.study(name)
        if checkpoint:
            self.checkpoint()
        self.studies.remove(st)
        self.sim_results.pop(name, None)
        self.dirty = True
        self._emit()

    def add_study_item(self, name: str, kind: str, itype: str, faces=None, checkpoint=True, object=None,
                       **params) -> dict:
        """kind: fixtures | loads | thermal. faces: face indices or ref specs on the study's object
        (see ref_on). Directions/points in world coordinates."""
        from ..sim.studies import face_refs
        st = self.study(name)
        spec = self.STUDY_ITEM_PARAMS.get(kind)
        if spec is None or itype not in spec:
            raise FeatureError(f"{kind}: type must be one of {list(spec or {})}")
        allowed = spec[itype]
        bad = [k for k in params if k not in allowed and k != "name"]
        if bad:
            raise FeatureError(f"{itype} takes {list(allowed) or 'no parameters'}, not {bad}")
        needs_faces = itype not in ("gravity", "centrifugal", "heat_volume")
        if needs_faces and not faces:
            raise FeatureError(f"{itype} needs faces")
        if kind == "thermal" and st["type"] not in ("thermal",):
            raise FeatureError("Thermal conditions go in a thermal study")
        if kind in ("fixtures", "loads") and st["type"] in ("thermal", "fatigue", "flow"):
            raise FeatureError(f"A {st['type']} study doesn't take {kind}")
        item = {"type": itype}
        on = st["object"]
        if object and object != on:
            if st["type"] != "assembly":
                raise FeatureError("Only assembly studies take faces on other objects")
            self.obj(object)
            on = object
        if st["type"] == "assembly":
            item["object"] = on
        try:
            if needs_faces:
                item["refs"] = face_refs(self, on, faces)
        except Exception as e:  # noqa: BLE001
            raise FeatureError(str(e)) from e
        for k, v in params.items():
            if k == "direction" and isinstance(v, str) and v != "normal":
                v = {"+X": [1, 0, 0], "-X": [-1, 0, 0], "+Y": [0, 1, 0], "-Y": [0, -1, 0], "+Z": [0, 0, 1],
                     "-Z": [0, 0, -1]}.get(v.upper(), None)
                if v is None:
                    raise FeatureError("direction: [x, y, z], '+X'…'-Z' or 'normal'")
            item[k] = v
        n = sum(1 for x in st[kind] if x["type"] == itype) + 1
        item.setdefault("name", f"{itype.replace('_', ' ').title()} {n}")
        if checkpoint:
            self.checkpoint()
        st[kind].append(item)
        self.dirty = True
        self._emit()
        return item

    def remove_study_item(self, name: str, kind: str, index: int, checkpoint=True):
        st = self.study(name)
        if not 0 <= index < len(st[kind]):
            raise FeatureError(f"No {kind} item {index}")
        if checkpoint:
            self.checkpoint()
        st[kind].pop(index)
        self.dirty = True
        self._emit()

    def run_study(self, name: str, progress=None, cancel=None):
        from ..sim import studies as S
        try:
            res = S.run(self, name, progress, cancel)
        except S.SimError as e:
            raise FeatureError(str(e)) from e
        self.sim_results[name] = res
        return res

    def study_result_fresh(self, name: str) -> bool:
        from ..sim.studies import fingerprint
        r = self.sim_results.get(name)
        return r is not None and getattr(r, "fingerprint", None) == fingerprint(self, self.study(name))

    # ------------------------------------------------------------ editing
    def get(self, name: str) -> Feature:
        for f in self.features:
            if f.name == name:
                return f
        raise FeatureError(f"No feature named '{name}'")

    def next_name(self, ftype: str) -> str:
        label = FEATURE_TYPES[ftype].LABEL.replace(" ", "")
        n = 1
        names = {f.name for f in self.features}
        while f"{label}{n}" in names:
            n += 1
        return f"{label}{n}"

    def add(self, ftype: str, params: dict | None = None, name: str | None = None,
            rebuild: bool = True, checkpoint: bool = True, layer: str | None = None,
            obj: str | None = None) -> Feature:
        if ftype not in FEATURE_TYPES:
            raise FeatureError(f"Unknown feature type '{ftype}'. Types: {list(FEATURE_TYPES)}")
        layer = layer or self.active_layer
        if not self.has_layer(layer):
            raise FeatureError(f"No layer named '{layer}'")
        if self.layer(layer).locked:
            raise FeatureError(f"Layer '{layer}' is locked — unlock it or pick another layer")
        if checkpoint:
            self.checkpoint()
        name = name or self.next_name(ftype)
        if name in self._fnames():
            raise FeatureError(f"Name '{name}' already used")
        obj = obj or self.active_object
        if not self.has_obj(obj):
            raise FeatureError(f"No object named '{obj}'")
        feat = FEATURE_TYPES[ftype](name, params, layer=layer, obj=obj)
        if self.rollback is not None:
            self.features.insert(self.rollback, feat)
            self.rollback += 1
            self._noted_add("feat", feat)
        else:
            self.features.append(feat)
            self._noted_add("feat", feat)
        self.dirty = True
        if rebuild:
            self.rebuild()
        return feat

    def edit(self, name: str, params: dict, rebuild=True, checkpoint=True):
        f = self.get(name)
        self._check_unlocked(f)
        if checkpoint:
            self.checkpoint()
        f.update(params)
        self.dirty = True
        if rebuild:
            self.rebuild()
        return f

    def rename(self, old: str, new: str):
        if not re.fullmatch(r"[A-Za-z_][\w\-]*", new):
            raise FeatureError("Names: letters, digits, _ and - only")
        if any(f.name == new for f in self.features):
            raise FeatureError(f"Name '{new}' already used")
        self._check_unlocked(self.get(old))
        self.checkpoint()
        f = self.get(old)
        f.name = new
        self._names_changed()
        for g in self.features:
            for p in g.PARAMS:
                if p.kind in ("sketch_ref", "feature_ref") and g.params.get(p.name) == old:
                    g.params[p.name] = new
        self.dirty = True
        self.rebuild()

    def delete(self, name: str, checkpoint=True):
        f = self.get(name)
        self._check_unlocked(f)
        users = [g.name for g in self.features if name in g.references()]
        if checkpoint:
            self.checkpoint()
        idx = self.features.index(f)
        self.features.remove(f)
        if self.rollback is not None and idx < self.rollback:
            self.rollback -= 1
        self.dirty = True
        self.rebuild()
        return users   # dependents now broken; caller may warn

    def set_suppressed(self, name: str, value: bool):
        self._check_unlocked(self.get(name))
        self.checkpoint()
        self.get(name).suppressed = value
        self.dirty = True
        self.rebuild()

    def move(self, name: str, new_index: int):
        f = self.get(name)
        self._check_unlocked(f)
        self.checkpoint()
        self.features.remove(f)
        self.features.insert(max(0, min(new_index, len(self.features))), f)
        self.dirty = True
        self.rebuild()

    def set_rollback(self, index: int | None):
        self.rollback = None if index is None or index >= len(self.features) else max(0, index)
        self.rebuild()

    def last_of(self, ftype: str) -> str | None:
        active = self.features[: self.rollback] if self.rollback is not None else self.features
        for f in reversed(active):
            if f.TYPE == ftype:
                return f.name
        return None

    # ------------------------------------------------------------ rebuild
    # rebuild cache: state after each feature, keyed by the hash of everything that produced it
    _BCACHE_MAX = 400

    @staticmethod
    def _ctx_copy(ctx):
        new = BuildContext()
        for k, v in vars(ctx).items():
            setattr(new, k, dict(v) if isinstance(v, dict) else list(v) if isinstance(v, list) else v)
        return new

    def rebuild(self) -> RebuildResult:
        """Evaluate every object's feature history. Each object resumes from the cached state of the longest
        unchanged prefix of its history, so an edit only re-runs the features from the edited one on, other
        objects aren't re-run at all, and undo/redo usually costs nothing."""
        import hashlib
        import json as _json
        from collections import OrderedDict
        t0 = time.perf_counter()
        cache = self.__dict__.setdefault("_bcache", OrderedDict())
        errors: dict[str, str] = {}
        eff = self.effective_features()
        active = eff[: self.rollback] if self.rollback is not None else eff
        off = {l.name for l in self.layers if not l.in_model}
        by_name = {f.name: f for f in eff}
        bodies, display, status, issues, cutlists = {}, {}, {}, {}, {}
        flats, threads, planes, surfaces = {}, {}, {}, {}
        makers_layers, makers_owners = {}, {}
        cfg_objs = self.config().get("objects", {}) if self.active_config != DEFAULT_CONFIG else {}
        for o in self.objects:
            o.suppressed = bool(cfg_objs.get(o.name, {}).get("suppressed", False))
        for o in self.objects:
            if o.instance_of or o.suppressed:
                continue
            mine = [f for f in active if f.obj == o.name]
            h = hashlib.sha1()
            h.update(f"{o.name}|{self.imported_step if o is self.objects[0] else ''}".encode())
            keys = []
            for f in mine:
                h.update(_json.dumps(f.to_dict(), sort_keys=True, default=str).encode())
                h.update(b"off" if f.layer in off else b"on")
                keys.append(h.hexdigest())
            start_i, ctx, oerr = 0, None, {}
            for i in range(len(mine) - 1, -1, -1):          # longest cached prefix
                hit = cache.get(keys[i])
                if hit is not None:
                    cache.move_to_end(keys[i])
                    ctx, oerr = self._ctx_copy(hit[0]), dict(hit[1])
                    start_i = i + 1
                    break
            if ctx is None:
                ctx = BuildContext()
                if self.imported_step and o is self.objects[0]:
                    try:
                        ctx.body = import_step(self.imported_step)
                    except Exception:  # noqa: BLE001
                        pass
            import datetime as _dt
            ctx.props = {"Name": o.name, "PartNo": getattr(o, "part_no", "") or o.name,
                         "Description": getattr(o, "description", "") or "",
                         "Material": getattr(self.material_of(o.name), "name", "") if hasattr(self, "material_of") else "",
                         "File": Path(self.path).stem if self.path else "", "Date": _dt.date.today().isoformat()}
            for i in range(start_i, len(mine)):
                f = mine[i]
                if not (f.suppressed or f.layer in off):
                    if any(r in oerr for r in f.references()):
                        oerr[f.name] = "Depends on a failed feature"
                    else:
                        foreign = [r for r in f.references() if r in by_name and by_name[r].obj != o.name]
                        if foreign:
                            oerr[f.name] = f"References {', '.join(foreign)} on another object"
                        else:
                            before = ctx.body
                            try:
                                f.build(ctx)
                                if ctx.body is not None and ctx.body is not before and f.TYPE != "import":
                                    if not _valid_cached(ctx.body):         # (imported geometry is taken as is)
                                        raise FeatureError("Kernel produced invalid geometry")
                            except Exception as e:  # noqa: BLE001 — kernel errors are varied
                                ctx.body = before
                                oerr[f.name] = str(e) or e.__class__.__name__
                cache[keys[i]] = (self._ctx_copy(ctx), dict(oerr))
                while len(cache) > max(self._BCACHE_MAX, 2 * len(self.features) + 64):
                    cache.popitem(last=False)
            errors.update(oerr)
            display.update(ctx.sketches)
            display.update(ctx.sketch_display)
            status.update(ctx.sketch_status)
            issues.update(ctx.sketch_issues)
            cutlists.update(ctx.cutlists)
            flats.update({k: {**v, "object": o.name} for k, v in ctx.flats.items()})
            threads.update({k: {**v, "object": o.name} for k, v in ctx.threads.items()})
            planes.update({k: (v, o.name) for k, v in ctx.planes.items()})
            surfaces.update({k: (v, o.name) for k, v in getattr(ctx, "surfaces", {}).items()})
            if ctx.body is not None:
                bodies[o.name] = ctx.body
                built = [f for f in mine if not (f.suppressed or f.layer in off)]
                makers_owners[o.name] = (lambda b=ctx.body, t=ctx.tools, n=[f.name for f in mine]:
                                         _owners_cached(b, t, n))
                lays = {f.layer for f in built}
                if len(lays) <= 1:                           # one layer: no attribution needed for colours
                    lay = next(iter(lays), DEFAULT)
                    makers_layers[o.name] = (lambda b=ctx.body, lay=lay: [lay] * nfaces(b))
                else:
                    makers_layers[o.name] = (lambda name=o.name, mk=makers_owners[o.name]:
                                             [by_name[w].layer if w in by_name else DEFAULT for w in mk()])
        for f in eff:                                        # sketch status shown on the real features
            src = getattr(f, "_origin", None)
            if src is not None and hasattr(f, "_status"):
                src._status = f._status
        for o in self.objects:                               # pattern / mirror instances share geometry
            if o.instance_of and not o.suppressed and not (self.has_obj(o.instance_of)
                                                           and self.obj(o.instance_of).suppressed):
                b = self.instance_body(o, bodies)
                if b is not None:
                    bodies[o.name] = b
                    if o.instance_of in makers_owners and not o.mirrored:
                        makers_owners[o.name] = makers_owners[o.instance_of]
                        makers_layers[o.name] = makers_layers[o.instance_of]
                    else:
                        makers_layers[o.name] = (lambda b=b: [DEFAULT] * len(b.faces()))
        if self.eq_errors:
            errors["Equations"] = "; ".join(f"{k}: {v}" for k, v in self.eq_errors.items())
        self.result = RebuildResult(body=None, sketches=display, errors=errors,
                                    ms=0.0, sketch_status=status, sketch_issues=issues,
                                    cutlists=cutlists, bodies=bodies,
                                    face_layers_by_obj=_LazyMap(makers_layers),
                                    feature_obj={f.name: f.obj for f in self.features},
                                    face_owner_by_obj=_LazyMap(makers_owners))
        self.result.flats, self.result.threads, self.result.planes = flats, threads, planes
        self.result.surfaces = surfaces
        self.solve_assembly()
        self._finish_result()
        self.result.ms = (time.perf_counter() - t0) * 1000
        self._emit()
        return self.result

    def _finish_result(self):
        """result.body: the one body (single object at the origin) or a compound of placed bodies."""
        r = self.result
        if len(self.objects) == 1 and not any(self.objects[0].pose):
            o = self.objects[0].name
            r.body = r.bodies.get(o)
            r.face_layers = r.face_layers_by_obj.get(o, [])
        else:
            placed = [b for b in self.placed_bodies().values()]
            if placed:
                from build123d import Compound
                r.body = Compound(placed) if len(placed) > 1 else placed[0]
            r.face_layers = []

    # ---------------------------------------------------- sketch support
    def context_before(self, name: str) -> BuildContext:
        """Evaluate the history up to (not including) `name`, without emitting."""
        self.get(name)
        ctx = BuildContext()
        off = {l.name for l in self.layers if not l.in_model}
        eff = self.effective_features()
        f = next(g for g in eff if g.name == name)
        for g in eff:
            if g is f:
                break
            if g.suppressed or g.layer in off or g.obj != f.obj:
                continue
            before = ctx.body
            try:
                g.build(ctx)
            except Exception:  # noqa: BLE001
                ctx.body = before
        return ctx

    def sketch_underlay(self, name: str, tol: float = 1e-3) -> list[list[tuple[float, float]]]:
        """Edges of the body (as it stands before `name`) lying in that sketch's
        plane, as plane-local polylines. Used as reference geometry while sketching."""
        from .features import resolve_plane
        f = self.get(name)
        ctx = self.context_before(name)
        if ctx.body is None:
            return []
        try:
            plane = resolve_plane(f.params["plane"], f.params["offset"], ctx)
        except Exception:  # noqa: BLE001
            return []
        out = []
        for e in ctx.body.edges():
            n = 2 if e.geom_type.name == "LINE" else 48
            pts = e.positions([i / (n - 1) for i in range(n)])
            loc = [plane.to_local_coords(p) for p in pts]
            if all(abs(p.Z) < tol for p in loc):
                out.append([(p.X, p.Y) for p in loc])
        return out

    # ------------------------------------------------------- evaluation
    def cut_list(self) -> list[dict]:
        """Grouped member list across all Frame features: section, length, qty."""
        rows: dict[tuple, dict] = {}
        for feat, members in self.result.cutlists.items():
            for m in members:
                key = (m["section"], m["long_point"], m["length"])
                r = rows.setdefault(key, {"frame": feat, "section": m["section"], "length": m["length"],
                                          "long_point": m["long_point"], "qty": 0})
                r["qty"] += 1
                if m.get("kg_per_m"):
                    r["kg_per_m"] = round(m["kg_per_m"], 4)
                    r["total_kg"] = round(m["kg_per_m"] * m["long_point"] / 1000 * r["qty"], 3)
        return sorted(rows.values(), key=lambda r: (r["section"], -r["long_point"]))

    # ---------------------------------------------------------- materials
    def material_of(self, obj: str | None = None) -> Material:
        name = self.material
        if obj is not None and self.has_obj(obj):
            o = self.obj(obj)
            if o.instance_of and not o.material and self.has_obj(o.instance_of):
                return self.material_of(o.instance_of)
            if o.material:
                name = o.material
        cm = self.config_material(obj) if hasattr(self, "active_config") else None
        if cm:
            name = cm
        try:
            return matlib.get(name, self.custom_materials)
        except KeyError:
            return matlib.get(matlib.DEFAULT)

    def all_materials(self) -> list[Material]:
        return list(matlib.LIBRARY) + [Material.from_dict(d) for d in self.custom_materials.values()]

    def set_material(self, name: str, obj: str | None = None, checkpoint=True):
        """Document material (obj None) or one object's own material ('' / None on an object = use the
        document's)."""
        if name:
            name = matlib.canonical(name)
            try:
                matlib.get(name, self.custom_materials)
            except KeyError:
                close = [m.name for m in matlib.search(name.split()[0] if name.split() else "",
                                                       custom=self.custom_materials)][:8]
                raise FeatureError(f"Unknown material '{name}'" + (f". Close: {close}" if close else ""))
        if checkpoint:
            self.checkpoint()
        if obj is None:
            if not name:
                raise FeatureError("Pick a material")
            self.material = name
        else:
            self.obj(obj).material = name or None
        self.dirty = True
        self._emit()

    def add_custom_material(self, props: dict, checkpoint=True) -> Material:
        base = matlib.get(props.get("based_on") or self.material, self.custom_materials).to_dict()
        base.update({k: v for k, v in props.items() if k in Material.__dataclass_fields__})
        base["custom"] = True
        base.setdefault("category", "Custom")
        if not base.get("name") or base["name"] in matlib.BY_NAME:
            raise FeatureError("A custom material needs a new name")
        m = Material.from_dict(base)
        if m.density <= 0 or m.E <= 0 or not (0 <= m.nu < 0.5):
            raise FeatureError("Density and E must be > 0, Poisson's ratio 0 … 0.5")
        if checkpoint:
            self.checkpoint()
        self.custom_materials[m.name] = m.to_dict()
        self.dirty = True
        self._emit()
        return m

    def mass_properties(self) -> dict:
        b = self.result.body
        if b is None:
            return {}
        bb = b.bounding_box()
        placed = self.placed_bodies()
        mass = sum(v.volume * self.material_of(n).density / 1e6 for n, v in placed.items()) if placed else 0.0
        vol = b.volume
        com = b.center()
        mats = sorted({self.material_of(n).name for n in placed}) or [self.material]
        return {
            "material": mats[0] if len(mats) == 1 else mats,
            "volume_mm3": round(vol, 3),
            "area_mm2": round(b.area, 3),
            "mass_kg": round(mass, 5),
            "bbox_mm": [round(bb.size.X, 3), round(bb.size.Y, 3), round(bb.size.Z, 3)],
            "bbox_min": [round(bb.min.X, 3), round(bb.min.Y, 3), round(bb.min.Z, 3)],
            "bbox_max": [round(bb.max.X, 3), round(bb.max.Y, 3), round(bb.max.Z, 3)],
            "center_of_mass": [round(com.X, 3), round(com.Y, 3), round(com.Z, 3)],
            "faces": len(b.faces()), "edges": len(b.edges()),
            **({"objects": {n: {"volume_mm3": round(v.volume, 3), "material": self.material_of(n).name,
                                "mass_kg": round(v.volume * self.material_of(n).density / 1e6, 5)}
                            for n, v in placed.items()}} if len(self.objects) > 1 else {}),
        }

    # ---------------------------------------------------------------- io
    def to_dict(self) -> dict:
        return {
            "format": "daedalus", "version": 1,
            "material": self.material, "custom_materials": self.custom_materials, "rollback": self.rollback,
            "imported_step": self.imported_step,
            "layers": [l.to_dict() for l in self.layers], "active_layer": self.active_layer,
            "objects": [o.to_dict() for o in self.objects], "active_object": self.active_object,
            "mates": [m.to_dict() for m in self.mates],
            "physics": self.physics.to_dict(),
            "studies": self.studies,
            "explode": self.explode,
            "drawings": [_strip_cache(dr.to_dict()) for dr in getattr(self, "drawings", [])],
            "export_options": getattr(self, "export_options", {}),
            "cam": getattr(self, "cam", {}),
            "features": [f.to_dict() for f in self.features],
            **self._v28_dict(),
        }

    def _load_dict(self, d: dict):
        if d.get("blobs"):
            from ..io.formats import BLOBS
            BLOBS.update(d["blobs"])
        self.layers = [Layer.from_dict(x) for x in d.get("layers", [])] or [Layer(DEFAULT)]
        if not self.has_layer(DEFAULT):
            self.layers.insert(0, Layer(DEFAULT))
        self.active_layer = d.get("active_layer", DEFAULT)
        if not self.has_layer(self.active_layer):
            self.active_layer = DEFAULT
        self.objects = [Obj.from_dict(x) for x in d.get("objects", [])] or [Obj("Object1", fixed=True)]
        self.active_object = d.get("active_object", self.objects[0].name)
        if not self.has_obj(self.active_object):
            self.active_object = self.objects[0].name
        self.mates = [Mate.from_dict(x) for x in d.get("mates", [])]
        self.physics = PhysicsSettings.from_dict(d.get("physics"))
        self.studies = [dict(x) for x in d.get("studies", [])]
        self.explode = dict(d.get("explode", {}))
        from ..drawing.model import Drawing as _Drawing
        self.drawings = [_Drawing.from_dict(x) for x in d.get("drawings", [])]
        self.export_options = dict(d.get("export_options", {}))
        self.cam = dict(d.get("cam", {}))
        self.sim_results = {k: v for k, v in getattr(self, "sim_results", {}).items()
                            if any(st["name"] == k for st in self.studies)}
        self.features = []
        for fd in d.get("features", []):
            cls = FEATURE_TYPES[fd["type"]]
            lay = fd.get("layer", DEFAULT)
            if not self.has_layer(lay):
                self.layers.append(Layer(lay, PALETTE[len(self.layers) % len(PALETTE)]))
            ob = fd.get("obj", self.objects[0].name)
            if not self.has_obj(ob):
                self.objects.append(Obj(ob))
            self.features.append(cls(fd["name"], fd.get("params"), fd.get("suppressed", False),
                                     layer=lay, obj=ob))
        self.custom_materials = dict(d.get("custom_materials", {}))
        self.material = matlib.canonical(d.get("material", self.material))
        for o in self.objects:
            if o.material:
                o.material = matlib.canonical(o.material)
        self.rollback = d.get("rollback")
        self.imported_step = d.get("imported_step")
        self._v28_load(d)

    def snapshot_to(self, path: str | Path, indent=1):
        """Write the document to `path` without touching its own path / dirty flag (autosave, backups)."""
        d = self.to_dict()
        from ..io.formats import blobs_for
        blobs = blobs_for(f.params.get("data") for f in self.features if f.TYPE == "import")
        if blobs:
            d["blobs"] = blobs
        Path(path).write_text(json.dumps(d, indent=indent))

    def save(self, path: str | Path):
        path = Path(path)
        self.snapshot_to(path)
        self.path, self.dirty = path, False
        self.exchange_files = []
        for fmt in (self.export_options or {}).get("auto_exchange", []) or []:
            try:                              # neutral copies for other CAD systems; never block the save
                if self.result.bodies:
                    self.exchange_files.append(self.export(path.with_suffix("." + fmt.lstrip("."))))
            except Exception:  # noqa: BLE001
                pass

    @classmethod
    def open(cls, path: str | Path) -> "Document":
        doc = cls()
        doc._load_dict(json.loads(Path(path).read_text()))
        doc.path = Path(path)
        doc.rebuild()
        return doc

    def set_export_options(self, **kw) -> dict:
        from ..io.formats import EXPORT_DEFAULTS, STEP_SCHEMAS
        bad = [k for k in kw if k not in EXPORT_DEFAULTS]
        if bad:
            raise FeatureError(f"Unknown export option(s) {bad}. Options: {sorted(EXPORT_DEFAULTS)}")
        if "step_schema" in kw and kw["step_schema"] not in STEP_SCHEMAS:
            raise FeatureError(f"step_schema must be one of {list(STEP_SCHEMAS)}")
        if "units" in kw and kw["units"] not in ("mm", "in"):
            raise FeatureError("units must be 'mm' or 'in'")
        if "auto_exchange" in kw:
            kw["auto_exchange"] = [str(f).lower().lstrip(".") for f in (kw["auto_exchange"] or [])]
            okf = {"step", "stp", "iges", "igs", "stl", "3mf", "obj", "glb", "brep", "ply", "wrl"}
            if any(f not in okf for f in kw["auto_exchange"]):
                raise FeatureError(f"auto_exchange formats: {sorted(okf)}")
        self.checkpoint()
        self.export_options = {**self.export_options, **kw}
        self.dirty = True
        self._emit()
        return {**EXPORT_DEFAULTS, **self.export_options}

    def export(self, path: str | Path, options: dict | None = None):
        """STEP (AP203/214/242 assembly with names/colours/placements), IGES, BREP, STL, OBJ, 3MF, PLY, VRML,
        glTF/GLB. options override the document's export options for this call."""
        from ..io.formats import ImportError_, export_assembly
        try:
            return export_assembly(self, path, options)
        except ImportError_ as e:
            raise FeatureError(str(e)) from e

    def add_flange(self, feature: str, point, length: float = 20.0, angle: float = 90.0, down: bool = False,
                   checkpoint=True, **opts) -> dict:
        """Edge flange on the sheet-metal outline edge nearest a world point (e.g. a picked side face)."""
        f = self.get(feature)
        if f.TYPE != "sheet_metal":
            raise FeatureError(f"{feature} isn't a sheet-metal feature")
        fl = self.result.flats.get(feature)
        if not fl or "edges" not in fl:
            raise FeatureError(f"{feature} isn't built")
        obj = fl.get("object") or f.obj
        M = np.linalg.inv(asm.pose_matrix(self.obj(obj).pose))
        lp = (M @ np.append(np.asarray(point, float), 1.0))[:3]
        from build123d import Vector
        q = fl["plane"].to_local_coords(Vector(*lp))
        best, bd = None, float("inf")
        for i, (a, b, line) in enumerate(fl["edges"]):
            if not line:
                continue
            a, b = np.array(a), np.array(b)
            ab = b - a
            tt = max(0.0, min(1.0, float(np.dot([q.X - a[0], q.Y - a[1]], ab) / max(ab @ ab, 1e-12))))
            dd = float(np.linalg.norm(a + ab * tt - [q.X, q.Y]))
            if dd < bd:
                best, bd = i, dd
        if best is None:
            raise FeatureError("No straight edge to flange")
        bad = [k for k in opts if k not in ("start_offset", "end_offset", "hem", "hem_length", "hem_gap",
                                             "length_ref", "position")]
        if bad:
            raise FeatureError(f"Unknown flange option(s) {bad}")
        flanges = [x for x in f.params["flanges"] if str(x.get("edge")) != str(best)]
        # an existing ±X/±Y flange on the same edge is replaced too
        from .sheetmetal import SIDES
        keep = []
        for x in flanges:
            if x.get("edge") in SIDES:
                ax = 0 if x["edge"][1] == "X" else 1
                sg = 1 if x["edge"][0] == "+" else -1
                lines = [i for i, e in enumerate(fl["edges"]) if e[2]]
                j = max(lines, key=lambda i: sg * (fl["edges"][i][0][ax] + fl["edges"][i][1][ax]) / 2
                        - 1e-3 * abs(fl["edges"][i][0][ax] - fl["edges"][i][1][ax]))
                if j == best:
                    continue
            keep.append(x)
        new = {"edge": best, "length": float(length), "angle": float(angle), "down": bool(down), **opts}
        self.edit(feature, {"flanges": keep + [new]}, checkpoint=checkpoint)
        return new

    def set_sheet_metal_state(self, state: str, feature: str | None = None, checkpoint=True) -> list[str]:
        """Fold ('folded') or flatten ('flat') one or every sheet-metal feature."""
        if state not in ("folded", "flat"):
            raise FeatureError("state: folded | flat")
        sm = ("sheet_metal", "convert_sheet_metal")
        names = [feature] if feature else [f.name for f in self.features if f.TYPE in sm]
        if not names:
            raise FeatureError("No sheet-metal features")
        if checkpoint:
            self.checkpoint()
        for n in names:
            f = self.get(n)
            if f.TYPE not in sm:
                raise FeatureError(f"{n} isn't sheet metal")
            f.update({"state": state})
        self.rebuild()
        return names

    def export_flat_dxf(self, feature: str, path: str | Path) -> dict:
        """Sheet-metal flat pattern → DXF (OUTLINE, BEND_UP / BEND_DOWN lines with notes)."""
        import ezdxf
        flat = self.result.flats.get(feature)
        if flat is None:
            raise FeatureError(f"'{feature}' isn't a built sheet-metal feature")
        from .features2 import flat_pattern_face
        faces = flat_pattern_face(flat)
        out = ezdxf.new("R2010", setup=True)
        out.header["$INSUNITS"] = 4
        for name, col in (("OUTLINE", 7), ("BEND_UP", 1), ("BEND_DOWN", 5), ("NOTES", 3)):
            out.layers.add(name, color=col)
        msp = out.modelspace()
        for w in [w_ for face in faces for w_ in (face.outer_wire(), *face.inner_wires())]:
            for e in w.edges():
                pts = [e.position_at(i / 16) for i in range(17)] if e.geom_type.name != "LINE" else \
                    [e.position_at(0), e.position_at(1)]
                msp.add_lwpolyline([(q.X, q.Y) for q in pts], dxfattribs={"layer": "OUTLINE"})
        for b in flat["bends"]:
            lay = "BEND_UP" if b["direction"] == "up" else "BEND_DOWN"
            (x0, y0), (x1, y1) = b["line"]
            msp.add_line((x0, y0), (x1, y1), dxfattribs={"layer": lay, "linetype": "DASHED"})
            msp.add_text(f"{b['direction'].upper()} {b['angle']:g}° R{b['radius']:g}",
                         dxfattribs={"layer": "NOTES", "height": 2.5}).set_placement(((x0 + x1) / 2, (y0 + y1) / 2))
        out.saveas(str(path))
        return {"size": flat["size"], "bends": len(flat["bends"]), "thickness": flat["thickness"]}

    def export_sketch_dxf(self, sketch: str, path: str | Path) -> int:
        from ..io.formats import ImportError_, export_sketch_dxf
        try:
            return export_sketch_dxf(self, sketch, path)
        except ImportError_ as e:
            raise FeatureError(str(e)) from e

    # ------------------------------------------------------------ library
    def insert_part(self, kind: str, size: str, length: float = 0.0, target: dict | None = None,
                    nut: bool = False, washer: bool = False, checkpoint=True) -> dict:
        """Standard part from the library as a new object. target = {"obj": name, "face": face index of a
        hole/bore (or a spec ref_on understands), "near": [x,y,z] world point by the end to use}: the part is
        placed in it and mated (concentric + face). Returns {"objects": [...], "mates": [...]}."""
        import numpy as np
        from . import standards as S
        from ..io.formats import shape_to_text, store_blob
        try:
            shape = S.model(kind, size, length)
        except (KeyError, ValueError) as e:
            raise FeatureError(str(e)) from e
        if checkpoint:
            self.checkpoint()
        made, mates = [], []
        mats = {"socket head cap screw": "Steel AISI 4140 (Q&T)", "hex bolt": "Steel AISI 1045 (CD)",
                "flat head screw": "Steel AISI 4140 (Q&T)", "hex nut": "Steel AISI 1045 (CD)",
                "washer": "Steel AISI 1018 (CD)", "ball bearing": "Steel AISI 4140 (Q&T)"}

        def add_obj(k, shp, desig, pose):
            nm = self._unique_obj_name(desig.split(" (")[0].replace("/", "-").replace('"', "in"))
            o = self.add_object(nm, "#9aa0a8" if k != "ball bearing" else "#c9ccd1", checkpoint=False,
                                activate=False, pose=pose)
            o.material, o.part_no = mats[k], desig
            o.unit_cost = S.unit_cost(k, size, length if k in ("socket head cap screw", "hex bolt",
                                                                "flat head screw") else 0.0)
            self.add("import", {"data": store_blob(shape_to_text(shp)), "source": desig, "format": "library"},
                     name=self._unique_feature_name(nm), rebuild=False, checkpoint=False, obj=o.name)
            made.append(o.name)
            return o

        if target is None:
            add_obj(kind, shape, S.designation(kind, size, length), [0.0] * 6)
            self.rebuild()
            return {"objects": made, "mates": []}
        tobj = target["obj"]
        body = self.result.bodies.get(tobj)
        if body is None:
            raise FeatureError(f"{tobj} has no body")
        spec = target.get("face")
        ref = self.ref_on(tobj, spec if spec is not None else {"axis_near": target.get("near_local", [0, 0, 0])})
        if ref["kind"] != "axis":
            raise FeatureError("Pick a hole or bore (cylindrical face)")
        f = list(body.faces())[ref["face"]]
        Mt = asm.pose_matrix(self.obj(tobj).pose)
        p0, d = np.array(ref["p"]), np.array(ref["d"]) / np.linalg.norm(ref["d"])
        hs = [float((np.array([v.X, v.Y, v.Z]) - p0) @ d) for v in f.vertices()] or [0.0]
        ends = [p0 + d * min(hs), p0 + d * max(hs)]
        near = target.get("near")
        if near is not None:
            nl = (np.linalg.inv(Mt) @ np.append(np.asarray(near, float), 1.0))[:3]
            which = int(np.argmin([np.linalg.norm(nl - e) for e in ends]))
        else:
            which = 1
        E, F = ends[which], ends[1 - which]
        e = (E - F) / (np.linalg.norm(E - F) or 1.0)          # out of the hole at the head end

        def pose_for(origin, zdir):
            z = zdir / np.linalg.norm(zdir)
            x = np.cross([0, 0, 1.0], z) if abs(z[2]) < 0.99 else np.array([1.0, 0, 0])
            x = x / np.linalg.norm(x)
            y = np.cross(z, x)
            R = np.column_stack([x, y, z])
            M = Mt @ np.vstack([np.column_stack([R, origin]), [0, 0, 0, 1]])
            return list(M[:3, 3]) + list(asm.rotvec(M[:3, :3]))

        def mate_pair(o_name, local_axis_pt, plane_n_local, plane_pt_local, face_pt, face_n, seat=None):
            """seat: (object, local plane point, local normal) the part sits on instead of the target face."""
            m1 = self.add_mate("concentric", {"obj": tobj, "ref": ref},
                               {"obj": o_name, "ref": {"kind": "axis", "p": list(local_axis_pt), "d": [0, 0, 1.0],
                                                       "r": 0.0, "face": None}}, checkpoint=False)
            if seat is None:
                side = {"obj": tobj, "ref": {"kind": "plane", "p": list(face_pt), "n": list(face_n), "face": None}}
            else:
                side = {"obj": seat[0], "ref": {"kind": "plane", "p": list(seat[1]), "n": list(seat[2]), "face": None}}
            m2 = self.add_mate("coincident", side,
                               {"obj": o_name, "ref": {"kind": "plane", "p": list(plane_pt_local),
                                                       "n": list(plane_n_local), "face": None}}, checkpoint=False)
            mates.extend([m1.name, m2.name])

        if kind == "ball bearing":
            dd, D, B = S.BEARINGS[size]
            mid = (E + F) / 2
            o = add_obj(kind, shape, S.designation(kind, size), pose_for(mid, e))
            self.rebuild()
            mate_pair(o.name, (0, 0, 0), (0, 0, 1.0), (0, 0, B / 2), mid + e * B / 2, e)
            return {"objects": made, "mates": mates}
        f_ = S.fastener(size)
        head_at = E
        seat = None
        if washer and kind in ("socket head cap screw", "hex bolt"):
            w = add_obj("washer", S.model("washer", size), S.designation("washer", size), pose_for(E, e))
            self.rebuild()
            mate_pair(w.name, (0, 0, 0), (0, 0, -1.0), (0, 0, 0), E, e)
            head_at = E + e * f_.washer_t
            seat = (w.name, (0, 0, f_.washer_t), (0, 0, 1.0))
        o = add_obj(kind, shape, S.designation(kind, size, length), pose_for(head_at, e))
        self.rebuild()
        if kind == "flat head screw":
            mate_pair(o.name, (0, 0, -1), (0, 0, 1.0), (0, 0, 0), E, e)
        else:
            mate_pair(o.name, (0, 0, -1), (0, 0, -1.0), (0, 0, 0), head_at, e, seat=seat)
        if nut and kind in ("socket head cap screw", "hex bolt"):
            n_ = add_obj("hex nut", S.model("hex nut", size), S.designation("hex nut", size), pose_for(F, -e))
            self.rebuild()
            mate_pair(n_.name, (0, 0, 1), (0, 0, -1.0), (0, 0, 0), F, -e)
        self.solve_assembly()
        self._emit()
        return {"objects": made, "mates": mates}

    # ---------------------------------------------------------------- BOM
    def bom(self) -> list[dict]:
        """Bill of materials: identical library parts grouped; one row per modelled object; mass and
        cost (raw material × 1.15 for modelled parts, catalogue price for library parts)."""
        rows: dict = {}
        for o in self.objects:
            b = self.result.bodies.get(o.name)
            if b is None:
                continue
            m = self.material_of(o.name)
            mass = b.volume * m.density / 1e6
            src = self.obj(o.instance_of) if o.instance_of and self.has_obj(o.instance_of) else None
            if src is not None and not o.part_no and not o.mirrored:
                o_key = src
            else:
                o_key = None
            if o_key is not None:
                key = ("obj", o_key.name)
                bb = b.bounding_box()
                desc = f"{bb.size.X:.0f} × {bb.size.Y:.0f} × {bb.size.Z:.0f} mm"
                unit = mass * m.price() * 1.15
            elif src is not None and o.mirrored and not o.part_no:
                key = ("obj", src.name + " (mirror)")
                bb = b.bounding_box()
                desc = f"Opposite hand of {src.name}"
                unit = mass * m.price() * 1.15
            elif o.part_no:
                key = ("lib", o.part_no)
                desc = o.description or "Purchased part"
                unit = float(o.unit_cost) if o.unit_cost is not None else (
                    mass * m.price() * 1.15 if (o.std or {}).get("kind") == "manufactured" else 0.0)
            else:
                key = ("obj", o.name)
                bb = b.bounding_box()
                desc = f"{bb.size.X:.0f} × {bb.size.Y:.0f} × {bb.size.Z:.0f} mm"
                flats = [v for v in self.result.flats.values() if v.get("object") == o.name]
                if flats:
                    fl = flats[0]
                    desc = f"Sheet t{fl['thickness']:g}, flat {fl['size'][0]:.0f} × {fl['size'][1]:.0f} mm, " \
                           f"{len(fl['bends'])} bends"
                frames = [f for f in self.features if f.obj == o.name and f.TYPE == "frame"]
                if frames:
                    desc = "Weldment: " + ", ".join(f.summary() for f in frames)
                unit = mass * m.price() * 1.15
            part = o.part_no or (key[1] if key[0] == "obj" else o.name)
            r = rows.setdefault(key, {"part": part, "description": desc, "material": m.name,
                                      "qty": 0, "unit_mass_kg": round(mass, 4), "unit_cost": round(unit, 2),
                                      "objects": []})
            if (o.source or {}).get("system") == "inventory":
                r["inventory_id"] = o.source.get("id")
            if o.std:
                r["standard"] = {"id": o.std.get("id"), "name": o.std.get("name"), "kind": o.std.get("kind"),
                                 "vendor": o.std.get("vendor"), "category": o.std.get("category")}
            r["qty"] += 1
            r["objects"].append(o.name)
            if o.group:
                r.setdefault("subassemblies", [])
                if o.group not in r["subassemblies"]:
                    r["subassemblies"].append(o.group)
        out = []
        for i, r in enumerate(sorted(rows.values(), key=lambda r: (r["part"] in [o.part_no for o in self.objects],
                                                                   r["part"])), start=1):
            r["item"] = i
            r["total_mass_kg"] = round(r["unit_mass_kg"] * r["qty"], 4)
            r["total_cost"] = round(r["unit_cost"] * r["qty"], 2)
            out.append(r)
        return out

    def export_bom(self, path: str | Path):
        """BOM (+ cut list, flat patterns) → .xlsx or .csv."""
        rows = self.bom()
        cols = ["item", "part", "description", "material", "qty", "unit_mass_kg", "total_mass_kg", "unit_cost",
                "total_cost"]
        path = Path(path)
        if path.suffix.lower() == ".csv":
            import csv
            with open(path, "w", newline="") as fh:
                w = csv.writer(fh)
                w.writerow(cols)
                for r in rows:
                    w.writerow([r[c] for c in cols])
                w.writerow([])
                w.writerow(["", "TOTAL", "", "", sum(r["qty"] for r in rows), "",
                            round(sum(r["total_mass_kg"] for r in rows), 3), "",
                            round(sum(r["total_cost"] for r in rows), 2)])
            return path
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError as e:
            raise FeatureError("Excel export needs openpyxl (pip install openpyxl) — or save as .csv") from e
        wb = Workbook()
        ws = wb.active
        ws.title = "BOM"
        head = ["Item", "Part", "Description", "Material", "Qty", "Unit mass (kg)", "Total mass (kg)",
                "Unit cost", "Total cost"]
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="2E3339")
        for r in rows:
            ws.append([r[c] for c in cols])
        n = len(rows) + 1
        ws.append(["", "TOTAL", "", "", f"=SUM(E2:E{n})", "", f"=SUM(G2:G{n})", "", f"=SUM(I2:I{n})"])
        for c in ws[n + 1]:
            c.font = Font(bold=True)
        for col, wdt in zip("ABCDEFGHI", (6, 34, 44, 28, 6, 14, 14, 11, 11)):
            ws.column_dimensions[col].width = wdt
        cl = self.cut_list()
        if cl:
            ws2 = wb.create_sheet("Cut list")
            ws2.append(["Frame", "Section", "Length (mm)", "Long point (mm)", "Qty"])
            for r in cl:
                ws2.append([r["frame"], r["section"], r["length"], r["long_point"], r["qty"]])
        if self.result.flats:
            ws3 = wb.create_sheet("Flat patterns")
            ws3.append(["Feature", "Object", "Thickness", "Flat X (mm)", "Flat Y (mm)", "Bends"])
            for k, v in self.result.flats.items():
                ws3.append([k, v.get("object"), v["thickness"], v["size"][0], v["size"][1], len(v["bends"])])
        wb.save(str(path))
        return path

    # ------------------------------------------------------- linked parts
    def insert_part_file(self, path: str | Path, link: bool = True, checkpoint=True, name: str | None = None) -> str:
        """Another .daedalus file as one object (all its objects placed together). Linked: re-read when the
        file changes (update_links / on open)."""
        import os
        from ..io.formats import shape_to_text, store_blob
        path = Path(path)
        src = _load_any(path)
        placed = list(src.placed_bodies().values())
        if not placed:
            raise FeatureError(f"{path.name} has no solid bodies")
        from build123d import Compound
        from ..io.formats import to_part
        shape = to_part(Compound(placed)) if len(placed) > 1 else placed[0]
        if checkpoint:
            self.checkpoint()
        o = self.add_object(self._unique_obj_name(name or path.stem), checkpoint=False, activate=False)
        o.material = src.material_of(src.objects[0].name).name if src.objects else None
        if o.material and o.material not in [m.name for m in self.all_materials()]:
            o.material = None
        self.add("import", {"data": store_blob(shape_to_text(shape)), "source": path.name, "format": path.suffix.lstrip(".").lower(),
                            "link": str(path.resolve()) if link else "", "mtime": os.path.getmtime(path)},
                 name=self._unique_feature_name(name or path.stem), rebuild=True, checkpoint=False, obj=o.name)
        return o.name

    def update_links(self) -> list[str]:
        """Re-read linked part files that changed on disk. Returns the features updated."""
        import os
        from build123d import Compound
        from ..io.formats import shape_to_text, store_blob, to_part
        done = []
        for f in self.features:
            link = f.params.get("link") if f.TYPE == "import" else None
            if not link or not os.path.exists(link):
                continue
            mt = os.path.getmtime(link)
            if mt <= float(f.params.get("mtime") or 0) + 1e-6:
                continue
            src = _load_any(Path(link))
            placed = list(src.placed_bodies().values())
            if not placed:
                continue
            shape = to_part(Compound(placed)) if len(placed) > 1 else placed[0]
            f.params["data"] = store_blob(shape_to_text(shape))
            f.params["mtime"] = mt
            done.append(f.name)
        if done:
            self.dirty = True
            self.rebuild()
        return done

    # ---------------------------------------------------------- drawings
    def drawing(self, name: str | None = None):
        if not self.drawings:
            raise FeatureError("No drawings yet — create one with new_drawing")
        if name is None:
            return self.drawings[0]
        for d in self.drawings:
            if d.name == name:
                return d
        raise FeatureError(f"No drawing named '{name}'. Drawings: {[d.name for d in self.drawings]}")

    def new_drawing(self, name: str | None = None, size: str = "A3", views=("front", "top", "right", "iso"),
                    objects=None, standard: str | None = None, units: str | None = None, landscape: bool = True,
                    title: dict | None = None, bom: bool | None = None, flat: str | None = None):
        """A drawing sheet with a standard view set laid out at an automatic scale."""
        from ..drawing import views as dv
        from ..drawing.model import SHEETS, Drawing, Sheet
        if size not in SHEETS:
            raise FeatureError(f"Sheet size must be one of {list(SHEETS)}")
        bad = [v for v in views if v not in ("front", "back", "top", "bottom", "right", "left", "iso", "dimetric",
                                              "trimetric")]
        if bad:
            raise FeatureError(f"Unknown view(s) {bad}")
        for o in objects or []:
            self.obj(o)
        name = name or f"Drawing{len(self.drawings) + 1}"
        if any(d.name == name for d in self.drawings):
            raise FeatureError(f"A drawing named '{name}' exists")
        self.checkpoint()
        dr = Drawing(name=name)
        if standard:
            dr.options.standard = standard
        if units:
            dr.options.units = units
        sh = Sheet(size=size, landscape=landscape, title=dict(title or {}))
        dr.sheets.append(sh)
        if flat:
            if flat not in self.result.flats:
                raise FeatureError(f"'{flat}' isn't a built sheet-metal feature")
            v = {"id": dv.new_id(), "kind": "flat", "feature": flat, "label": "FLAT PATTERN"}
            dv.compute(self, v, dr.options)
            x0, y0, x1, y1 = v["_bbox"]
            W, H = sh.dims()
            sh.scale = dv.nice_scale(x1 - x0, y1 - y0, W - 60, H - 90)
            from ..drawing.model import scale_factor
            v.update(x=W * 0.42 - (x0 + x1) / 2 * scale_factor(sh.scale), y=H * 0.58 - (y0 + y1) / 2 *
                     scale_factor(sh.scale), scale=scale_factor(sh.scale))
            sh.views.append(v)
        elif views:
            placed = dv.standard_views(self, dr, sh, tuple(views), objects=list(objects or []))
            multi = len([o for o in self.objects if self.result.bodies.get(o.name) is not None]) > 1
            if placed and (bom if bom is not None else multi):
                iso = next((v for v in placed if v["dir"] in ("iso", "dimetric", "trimetric")), placed[0])
                iso["balloons"] = True
                iso["bom"] = True
        self.drawings.append(dr)
        dv.regenerate(self, dr)
        self.dirty = True
        self._emit()
        return dr

    def add_drawing_view(self, drawing: str | None = None, kind: str = "projected", dir: str = "front",
                         x: float | None = None, y: float | None = None, scale: float | None = None,
                         parent: str | None = None, line=None, flip: bool = False, center=None, radius=None,
                         label: str = "", objects=None, sheet: str | None = None, exploded=False, hidden=None,
                         feature: str | None = None, breaks=None):
        """Add a view: projected/iso (dir), section (line [[x,y],[x,y]] in the parent view), detail
        (center [x,y] + radius in the parent view) or flat (sheet-metal feature)."""
        from ..drawing import views as dv
        from ..drawing.model import VIEW_DIRS, scale_factor
        dr = self.drawing(drawing)
        sh = dr.sheet(sheet)
        if kind not in ("projected", "base", "iso", "section", "detail", "flat", "auxiliary"):
            raise FeatureError("kind: projected | base | iso | section | detail | flat | auxiliary")
        byid = {v["id"]: v for v in sh.views}
        bylabel = {v.get("label") or v["id"]: v for v in sh.views}
        par = None
        if parent:
            par = byid.get(parent) or bylabel.get(parent) or next((v for v in sh.views if v.get("dir") == parent),
                                                                 None)
            if par is None:
                raise FeatureError(f"No view '{parent}' on {sh.name}")
        elif kind in ("section", "detail", "auxiliary"):
            par = next((v for v in sh.views if v.get("kind") in ("base", "projected")), None)
            if par is None:
                raise FeatureError("Section/detail/auxiliary views need a parent view")
        if par is not None and "_prims" not in par:
            dv.regenerate(self, dr)
        W, H = sh.dims()
        s_sheet = scale_factor(sh.scale)
        v = {"id": dv.new_id(), "kind": kind, "objects": list(objects or (par or {}).get("objects", [])),
             "label": label, "exploded": exploded}
        if hidden is not None:
            v["hidden"] = hidden
        letters = [vv.get("section", {}).get("name") for vv in sh.views] + \
                  [vv.get("detail", {}).get("name") for vv in sh.views]
        nxt = next(ch for ch in "ABCDEFGHJKLMNPRSTUVWXYZ" if ch not in letters)
        if kind in ("projected", "base", "iso"):
            if dir not in VIEW_DIRS:
                raise FeatureError(f"dir must be one of {list(VIEW_DIRS)}")
            v["dir"] = dir
            if kind == "iso" or dir in ("iso", "dimetric", "trimetric"):
                v["kind"], v["hidden"] = "iso", v.get("hidden", False)
            if par is not None:
                v["parent"], v["center"] = par["id"], par.get("center")
        elif kind == "section":
            if not line or len(line) != 2:
                raise FeatureError("A section needs line=[[x1,y1],[x2,y2]] in the parent view's mm")
            d, u, r = (np.asarray(q, float) for q in par["_basis"])
            c = np.asarray(par["center"], float)
            a3 = c + r * line[0][0] + u * line[0][1]
            b3 = c + r * line[1][0] + u * line[1][1]
            L3 = b3 - a3
            if np.linalg.norm(L3) < 1e-9:
                raise FeatureError("Section line has zero length")
            n = np.cross(L3, d)
            n = n / np.linalg.norm(n) * (-1 if flip else 1)
            v.update(parent=par["id"], dir="custom", vector=list(n), up=list(d if abs(n @ d) < 0.99 else u),
                     section={"point": list(a3), "normal": list(n), "line": [list(line[0]), list(line[1])],
                              "name": nxt}, center=par.get("center"))
            v["up"] = list(np.cross(n, L3 / np.linalg.norm(L3)))
        elif kind == "auxiliary":
            if not line or len(line) != 2:
                raise FeatureError("An auxiliary view needs line=[[x1,y1],[x2,y2]] along an edge in the parent view")
            d, u, r = (np.asarray(q, float) for q in par["_basis"])
            e2 = np.asarray(line[1], float) - np.asarray(line[0], float)
            if np.linalg.norm(e2) < 1e-9:
                raise FeatureError("Auxiliary edge has zero length")
            e2 = e2 / np.linalg.norm(e2)
            n2 = np.array([-e2[1], e2[0]]) * (-1 if flip else 1)
            vec = r * n2[0] + u * n2[1]
            v.update(parent=par["id"], dir="custom", vector=list(vec), up=list(d), center=par.get("center"),
                     aux={"line": [list(line[0]), list(line[1])], "name": nxt})
            v["label"] = label or f"VIEW {nxt}"
        elif kind == "detail":
            if center is None or radius is None:
                raise FeatureError("A detail needs center=[x,y] and radius in the parent view's mm")
            v.update(parent=par["id"], detail={"c": list(center), "r": float(radius), "name": nxt})
            scale = scale or par.get("scale", s_sheet) * 2
        elif kind == "flat":
            if feature not in self.result.flats:
                raise FeatureError(f"'{feature}' isn't a built sheet-metal feature")
            v.update(feature=feature, label=label or "FLAT PATTERN")
        v["scale"] = float(scale or (par or {}).get("scale", s_sheet))
        v["_scale"] = v["scale"]
        if breaks:
            v["breaks"] = [dict(b) for b in breaks]
        self.checkpoint()
        sh.views.append(v)
        dv.regenerate(self, dr)
        if x is None or y is None:                  # first free spot on a coarse grid, right of / below the parent
            x0, y0, x1, y1 = v.get("_bbox", (0, 0, 0, 0))
            vw, vh = (x1 - x0) * v["scale"], (y1 - y0) * v["scale"]
            taken = []
            for o in sh.views:
                if o is v or "x" not in o:
                    continue
                a, b, c_, d_ = o.get("_bbox", (0, 0, 0, 0))
                taken.append((o["x"] + a * o["scale"] - 8, o["y"] + b * o["scale"] - 14, o["x"] + c_ * o["scale"] + 8,
                              o["y"] + d_ * o["scale"] + 8))
            taken.append((W - 200, 0, W, 60))             # title block
            best = None
            for gy in np.arange(H - 25 - vh / 2, 25 + vh / 2, -5.0):
                for gx in np.arange(25 + vw / 2, W - 25 - vw / 2, 5.0):
                    box = (gx - vw / 2, gy - vh / 2, gx + vw / 2, gy + vh / 2)
                    if all(box[2] < t[0] or box[0] > t[2] or box[3] < t[1] or box[1] > t[3] for t in taken):
                        best = (gx, gy)
                        break
                if best:
                    break
            best = best or (W * 0.7, H * 0.5)
            x = best[0] - (x0 + x1) / 2 * v["scale"]
            y = best[1] - (y0 + y1) / 2 * v["scale"]
        v["x"], v["y"] = float(x), float(y)
        dv.regenerate(self, dr)
        self.dirty = True
        self._emit()
        return v

    def set_drawing(self, drawing: str | None = None, sheet: str | None = None, options: dict | None = None,
                    title: dict | None = None, size: str | None = None, landscape: bool | None = None,
                    scale=None, rename: str | None = None):
        from ..drawing import views as dv
        from ..drawing.model import SHEETS, Options
        dr = self.drawing(drawing)
        sh = dr.sheet(sheet)
        if options:
            bad = [k for k in options if k not in Options.__dataclass_fields__]
            if bad:
                raise FeatureError(f"Unknown drawing option(s) {bad}. Options: {list(Options.__dataclass_fields__)}")
        if size and size not in SHEETS:
            raise FeatureError(f"Sheet size must be one of {list(SHEETS)}")
        self.checkpoint()
        if options:
            for k, val in options.items():
                setattr(dr.options, k, type(getattr(dr.options, k))(val))
        if title:
            sh.title.update({k: str(v) for k, v in title.items()})
        if size:
            sh.size = size
        if landscape is not None:
            sh.landscape = bool(landscape)
        if scale:
            a, b = (scale if isinstance(scale, (list, tuple)) else [float(q) for q in str(scale).split(":")])
            old = sh.scale[0] / sh.scale[1]
            sh.scale = [a, b]
            for v in sh.views:
                if abs(v.get("scale", old) - old) < 1e-9 or v.get("kind") != "detail":
                    v["scale"] = a / b if v.get("kind") != "detail" else v["scale"] * (a / b) / old
        if rename:
            dr.name = rename
        dv.regenerate(self, dr)
        self.dirty = True
        self._emit()
        return dr

    def set_drawing_view(self, view: str, drawing: str | None = None, sheet: str | None = None, breaks=None,
                         label: str | None = None, scale: float | None = None, x: float | None = None,
                         y: float | None = None, hidden: bool | None = None):
        """Change a view: breaks [{"axis": "x"|"y", "a": lo, "b": hi}] in view mm (broken view; [] = none),
        label, scale, position, hidden lines."""
        from ..drawing import views as dv
        dr = self.drawing(drawing)
        sh = dr.sheet(sheet)
        v = next((q for q in sh.views if q["id"] == view or (q.get("label") or "") == view or q.get("dir") == view),
                 None)
        if v is None:
            raise FeatureError(f"No view '{view}' on {sh.name}")
        self.checkpoint()
        if breaks is not None:
            for b in breaks:
                if b.get("axis", "x") not in ("x", "y") or "a" not in b or "b" not in b:
                    raise FeatureError("breaks: [{axis: x|y, a, b}]")
            v["breaks"] = [dict(b) for b in breaks]
        if label is not None:
            v["label"] = label
        if scale:
            v["scale"] = float(scale)
        if x is not None:
            v["x"] = float(x)
        if y is not None:
            v["y"] = float(y)
        if hidden is not None:
            v["hidden"] = bool(hidden)
        dv.regenerate(self, dr)
        self.dirty = True
        self._emit()
        return v

    def add_revision(self, description: str, drawing: str | None = None, sheet: str | None = None,
                     rev: str | None = None, date: str | None = None, approved: str = "", zone: str = "") -> dict:
        """Add a row to the sheet's revision table (rev letter auto-increments); the title block REV follows."""
        import datetime as _dt

        from ..drawing.annot import next_revision
        dr = self.drawing(drawing)
        sh = dr.sheet(sheet)
        self.checkpoint()
        row = {"rev": rev or next_revision(sh.revisions), "description": str(description),
               "date": date or _dt.date.today().isoformat(), "approved": approved, "zone": zone}
        sh.revisions.append(row)
        self.dirty = True
        self._emit()
        return row

    def add_ordinate_dims(self, origin, points, axis: str = "x", drawing: str | None = None,
                          sheet: str | None = None, offset: float = 10.0) -> list[dict]:
        """Ordinate dimensions (sheet mm) from an origin point to each point, aligned on one line."""
        from ..drawing.annot import ordinate_set
        from ..drawing.model import scale_factor
        if axis not in ("x", "y"):
            raise FeatureError("axis: x | y")
        dr = self.drawing(drawing)
        sh = dr.sheet(sheet)
        items = ordinate_set(origin, points, axis, offset, unit=1.0 / scale_factor(sh.scale) if sh.dims()[0] else None)
        self.checkpoint()
        sh.items.extend(items)
        self.dirty = True
        self._emit()
        return items

    def add_drawing_item(self, item: dict, drawing: str | None = None, sheet: str | None = None):
        """A 2D entity in sheet mm (see daedalus.drawing.model for the primitive schema)."""
        if item.get("t") not in ("line", "poly", "circle", "arc", "text", "hatch", "dim", "balloon", "gdt", "datum",
                                 "finish"):
            raise FeatureError("item t: line | poly | circle | arc | text | hatch | dim | balloon | gdt | datum | "
                               "finish")
        if item.get("t") == "gdt":
            from ..drawing.annot import GDT
            if item.get("symbol", "position") not in GDT:
                raise FeatureError(f"GD&T symbol must be one of {list(GDT)}")
        dr = self.drawing(drawing)
        self.checkpoint()
        dr.sheet(sheet).items.append(dict(item))
        self.dirty = True
        self._emit()
        return item

    def delete_drawing(self, name: str):
        dr = self.drawing(name)
        self.checkpoint()
        self.drawings.remove(dr)
        self._emit()

    def regenerate_drawings(self):
        from ..drawing import views as dv
        for dr in self.drawings:
            if dr.linked:
                dv.regenerate(self, dr)

    def export_drawing(self, path: str | Path, drawing: str | None = None, sheet: str | None = None, **kw) -> str:
        """PDF (all sheets), SVG, PNG, DXF or DWG (DWG needs the ODA File Converter)."""
        from ..drawing import dxfio
        from ..drawing import views as dv
        dr = self.drawing(drawing)
        if dr.linked:
            dv.regenerate(self, dr)
        try:
            return dxfio.save_2d(self, dr, path, dr.sheet(sheet) if sheet else None, **kw)
        except dxfio.DwgUnavailable as e:
            raise FeatureError(str(e)) from e
        except ValueError as e:
            raise FeatureError(str(e)) from e

    def open_2d(self, path: str | Path):
        """Open a DXF/DWG/SVG/HPGL file as an editable 2D drawing."""
        from ..drawing import dxfio
        try:
            dr = dxfio.read_2d(path)
        except dxfio.DwgUnavailable as e:
            raise FeatureError(str(e)) from e
        except ValueError as e:
            raise FeatureError(str(e)) from e
        base, i = dr.name, 2
        while any(d.name == dr.name for d in self.drawings):
            dr.name = f"{base} ({i})"
            i += 1
        self.checkpoint()
        self.drawings.append(dr)
        self.dirty = True
        self._emit()
        return dr

    # --------------------------------------------------------------- CAM
    def cam_setup(self, **kw) -> dict:
        """object, margin [x,y,z], origin corner|center, post fanuc|haas|linuxcnc|grbl, units mm|in, max_rpm,
        safe_z, clearance_z, coolant, program."""
        from ..manufacturing import cam
        base = {**cam.default_setup(), **(self.cam or {})}
        bad = [k for k in kw if k not in base or k in ("tools", "ops")]
        if bad:
            raise FeatureError(f"Unknown CAM setting(s) {bad}")
        if kw.get("object"):
            self.obj(kw["object"])
        if "post" in kw and kw["post"] not in cam.POSTS:
            raise FeatureError(f"post must be one of {list(cam.POSTS)}")
        if "origin" in kw and kw["origin"] not in ("corner", "center"):
            raise FeatureError("origin: corner | center")
        if "orient" in kw and kw["orient"] not in cam.ORIENT:
            raise FeatureError(f"orient must be one of {list(cam.ORIENT)}")
        self.checkpoint()
        self.cam = {**base, **kw}
        self.dirty = True
        self._emit()
        return {k: v for k, v in self.cam.items() if k not in ("tools", "ops")}

    def cam_add_tool(self, kind: str, diameter: float, flutes: int = 3, number: int | None = None, name: str = "",
                     rpm: float = 0.0, feed: float = 0.0, pitch: float = 0.0) -> dict:
        from ..manufacturing import cam
        if kind not in ("flat", "face", "drill", "tap", "spot", "ball"):
            raise FeatureError("tool kind: flat | face | drill | tap | spot | ball")
        if diameter <= 0:
            raise FeatureError("diameter must be > 0")
        setup = {**cam.default_setup(), **(self.cam or {})}
        nums = [t["number"] for t in setup["tools"]]
        number = int(number or (max(nums, default=0) + 1))
        self.checkpoint()
        setup["tools"] = [t for t in setup["tools"] if t["number"] != number] + [
            {"number": number, "kind": kind, "diameter": float(diameter), "flutes": int(flutes), "name": name,
             "rpm": float(rpm), "feed": float(feed), "pitch": float(pitch)}]
        setup["tools"].sort(key=lambda t: t["number"])
        self.cam = setup
        self.dirty = True
        self._emit()
        return setup["tools"][[t["number"] for t in setup["tools"]].index(number)]

    def cam_add_op(self, type: str, tool: int | None = None, **params) -> list[dict]:
        """facing | pocket | profile | drill. Without a tool one is chosen (drill: one op per hole size, adding
        drills/taps to the tool table as needed)."""
        from ..manufacturing import cam
        if type not in cam.OP_TYPES:
            raise FeatureError(f"type must be one of {list(cam.OP_TYPES)}")
        setup = {**cam.default_setup(), **(self.cam or {})}
        setup["tools"] = [dict(t) for t in setup["tools"]]
        setup["ops"] = [dict(o) for o in setup["ops"]]
        tools = setup["tools"]
        made = []
        self.checkpoint()

        def new_tool(kind, dia, name, pitch=0.0):
            for t in tools:
                if t["kind"] == kind and abs(t["diameter"] - dia) < 1e-6:
                    return t["number"]
            n = max([t["number"] for t in tools], default=0) + 1
            tools.append({"number": n, "kind": kind, "diameter": round(dia, 4), "flutes": 2, "name": name,
                          "pitch": pitch})
            return n

        if tool is not None and not any(t["number"] == int(tool) for t in tools):
            raise FeatureError(f"No tool T{tool}. Tools: {[t['number'] for t in tools]}")
        if type == "drill" and tool is None:
            name = setup.get("object") or self.objects[0].name
            body = self.placed_bodies().get(name)
            if body is None:
                raise FeatureError(f"{name} has no body")
            holes = cam.vertical_holes(body)
            if not holes:
                raise FeatureError("No vertical holes to drill")
            threads = {(round(t["axis"][0][0], 3), round(t["axis"][0][1], 3)): t
                       for t in self.result.threads.values()}
            for dia in sorted({round(2 * h["r"], 3) for h in holes}):
                n = new_tool("drill", dia, f"Ø{dia:g} drill")
                op = {"type": "drill", "tool": n, "diameter": dia, **params}
                setup["ops"].append(op)
                made.append(op)
                tapped = [threads[k] for k in threads if any(abs(h["x"] - k[0]) < 1e-3 and abs(h["y"] - k[1]) < 1e-3
                                                              and abs(2 * h["r"] - dia) < 1e-3 for h in holes)]
                if tapped:
                    th = tapped[0]
                    try:
                        tn = new_tool("tap", float(th["major"]), f"{th['size']} tap", float(th["pitch"]))
                        op2 = {"type": "drill", "tool": tn, "diameter": dia, **params}
                        setup["ops"].append(op2)
                        made.append(op2)
                    except Exception:  # noqa: BLE001
                        pass
        else:
            if tool is None:
                cands = {"facing": [t for t in tools if t["kind"] == "face"],
                         "pocket": [t for t in tools if t["kind"] == "flat"],
                         "profile": [t for t in tools if t["kind"] == "flat"]}[type]
                cands = cands or [t for t in tools if t["kind"] in ("flat", "face")]
                if not cands:
                    raise FeatureError("Add a milling tool first")
                pick = max(cands, key=lambda t: t["diameter"])
                if type == "pocket":                     # largest tool that fits every pocket
                    body = self.placed_bodies().get(setup.get("object") or self.objects[0].name)
                    for t in sorted(cands, key=lambda t: -t["diameter"]):
                        trial = {**setup, "ops": [{"type": "pocket", "tool": t["number"], **params}]}
                        try:
                            r = cam.plan(self, trial)
                            if not any("doesn't fit" in n_ for n_ in r["ops"][0].notes):
                                pick = t
                                break
                        except Exception:  # noqa: BLE001
                            continue
                    del body
                tool = pick["number"]
            op = {"type": type, "tool": int(tool), **params}
            setup["ops"].append(op)
            made.append(op)
        self.cam = setup
        self.dirty = True
        self._emit()
        return made

    def cam_edit_op(self, index: int, **params) -> dict:
        """Change an operation's settings. Drilling: cycle (auto|G81|G82|G83|G73|G84|G74|G85|G86|G89|G76),
        peck (Q, mm), dwell (s), retract ('initial' = G98 | 'r' = G99), r_clear (R above the hole top, mm),
        shift (G76 lift-off, mm), breakthrough (mm), diameter_tolerance; any op: name, tool, suppressed."""
        from ..manufacturing import cam
        setup = {**cam.default_setup(), **(self.cam or {})}
        ops = [dict(o) for o in setup["ops"]]
        if not 0 <= index < len(ops):
            raise FeatureError(f"No operation {index}; there are {len(ops)}")
        allowed = set(cam.CYCLE_PARAMS) | {"name", "tool", "suppressed", "stepdown", "stepover", "depth"}
        bad = [k for k in params if k not in allowed]
        if bad:
            raise FeatureError(f"Unknown setting(s) {bad}. Allowed: {sorted(allowed)}")
        if "cycle" in params:
            c = str(params["cycle"])
            c = "auto" if c.lower() == "auto" else c.upper()
            if c not in cam.CYCLES:
                raise FeatureError(f"cycle must be one of {list(cam.CYCLES)}")
            params["cycle"] = c
        if "retract" in params and str(params["retract"]).lower() not in ("initial", "r", "g98", "g99"):
            raise FeatureError("retract must be 'initial' (G98) or 'r' (G99)")
        if "tool" in params and not any(t["number"] == int(params["tool"]) for t in setup["tools"]):
            raise FeatureError(f"No tool T{params['tool']}")
        self.checkpoint()
        ops[index].update(params)
        setup["ops"] = ops
        self.cam = setup
        self.dirty = True
        self._emit()
        return ops[index]

    def cam_remove_op(self, index: int):
        ops = (self.cam or {}).get("ops", [])
        if not 0 <= index < len(ops):
            raise FeatureError(f"No operation {index + 1}")
        self.checkpoint()
        self.cam = {**self.cam, "ops": [o for i, o in enumerate(ops) if i != index]}
        self.dirty = True
        self._emit()

    def cam_plan(self):
        from ..manufacturing import cam
        if not (self.cam or {}).get("ops"):
            raise FeatureError("No CAM operations — add facing / pocket / profile / drill first")
        try:
            self.last_cam = cam.plan(self)
        except ValueError as e:
            raise FeatureError(str(e)) from e
        return cam.summary(self.last_cam)

    def export_gcode(self, path: str | Path, post: str | None = None) -> dict:
        from ..manufacturing import cam
        self.cam_plan()
        try:
            text = cam.gcode(self.last_cam, post)
        except ValueError as e:
            raise FeatureError(str(e)) from e
        Path(path).write_text(text)
        s = cam.summary(self.last_cam)
        s["lines"] = text.count("\n")
        s["path"] = str(path)
        return s

    def check_manufacturability(self, check: str = "draft", object: str | None = None, direction=(0, 0, 1),
                                min_draft: float = 1.0, min_wall: float = 1.0, max_overhang: float = 45.0) -> dict:
        """draft (mould pull + undercuts) | thickness | overhang (3D printing, direction = build direction)."""
        from ..manufacturing import printability as pr
        name = object or self.active_object
        body = self.placed_bodies().get(name)
        if body is None:
            raise FeatureError(f"{name} has no body")
        try:
            r = pr.analyze(body, check, direction, min_draft=min_draft, min_wall=min_wall, max_overhang=max_overhang)
        except ValueError as e:
            raise FeatureError(str(e)) from e
        if check == "overhang":
            r["summary"]["best_build_dir"] = pr.best_orientation(body, max_overhang)
        r["object"] = name
        self.last_check = r
        return r["summary"]

    # ----------------------------------------------------------- explode
    def auto_explode(self, factor: float = 1.0, checkpoint=True) -> dict:
        """Exploded-view offsets: each movable object pushed away from the assembly centre (fasteners
        along their own axis)."""
        import numpy as np
        placed = self.placed_bodies()
        if not placed:
            return {}
        cs = {n: np.array(b.center().to_tuple() if hasattr(b.center(), "to_tuple") else
                          (b.center().X, b.center().Y, b.center().Z)) for n, b in placed.items()}
        allc = np.mean(list(cs.values()), axis=0)
        size = max(b.bounding_box().diagonal for b in placed.values())
        if checkpoint:
            self.checkpoint()
        out = {}
        for o in self.objects:
            if o.fixed or o.name not in cs:
                continue
            if o.part_no:
                d = asm.rotmat(o.pose[3:6]) @ np.array([0, 0, 1.0])      # out along the fastener's own axis
            else:
                d = cs[o.name] - allc
                if np.linalg.norm(d) < 1e-6:
                    d = np.array([0, 0, 1.0])
            d = d / np.linalg.norm(d)
            out[o.name] = [float(v) for v in d * size * 0.6 * factor]
        self.explode = out
        self.exploded = True
        self.dirty = True
        self._emit()
        return out

    def display_poses(self) -> dict:
        poses = {o.name: list(o.pose) for o in self.objects}
        if getattr(self, "exploded", False):
            for n, off in getattr(self, "explode", {}).items():
                if n in poses:
                    poses[n] = [float(poses[n][i] + off[i]) for i in range(3)] + list(poses[n][3:])
        return poses

    # ------------------------------------------------------------- import
    def _unique_obj_name(self, base: str) -> str:
        base = re.sub(r"[/\\]", "_", (base or "Part").strip())[:36].strip() or "Part"
        name, n = base, 2
        while self.has_obj(name) or not valid_name(name):
            name = f"{base}-{n}" if valid_name(f"{base}-{n}") else f"Part{n}"
            n += 1
        return name

    def _unique_feature_name(self, base: str) -> str:
        base = re.sub(r"[^\w\-]", "_", base or "Import")
        if not re.match(r"[A-Za-z_]", base):
            base = "_" + base
        base = base[:40]
        names = self._fnames()
        name, n = base, 2
        while name in names:
            name, n = f"{base}_{n}", n + 1
        return name

    def import_file(self, path: str | Path, checkpoint=True, plane: str = "XY", progress=None) -> dict:
        """Bring a file in. Solids: one object per part (name, colour, placement kept); a single
        part goes into the active object when that object is still empty. DXF: a new editable
        sketch in the active object. SolidWorks bridge output also brings its mates.
        Returns a summary {kind, objects, features, mates, notes}."""
        from ..io.formats import ImportError_, import_file
        path = Path(path)
        try:
            res = import_file(path, progress)
        except ImportError_ as e:
            raise FeatureError(str(e)) from e
        return self.import_result(res, path, checkpoint=checkpoint, plane=plane, progress=progress)

    def import_result(self, res, path, checkpoint=True, plane: str = "XY", progress=None) -> dict:
        """Apply an already-read ImportResult (e.g. read in a background process). Parts may carry
        `blob` (serialized geometry) instead of a shape."""
        from ..io.formats import shape_to_text, store_blob
        from ..io.formats import _cache as fcache
        from ..io.formats import register_geometry_key
        path = Path(path)
        blob_of: dict = {}                    # STEP prototype → blob (instances share one geometry)
        if checkpoint:
            self.checkpoint()
        summary = {"kind": res.kind, "objects": [], "features": [], "mates": [], "notes": list(res.notes)}
        try:
            if res.sketch is not None:
                name = self._unique_feature_name(path.stem)
                self.add("sketch", {"plane": plane, "geometry": res.sketch}, name=name,
                         rebuild=False, checkpoint=False)
                summary["features"].append(name)
                n = len(res.sketch.get("curves", []))
                if n > 150:
                    summary["notes"].append(f"{n} curves — dimensions/relations are up to you; "
                                            "large sketches solve slower")
            fmt = res.kind if res.kind not in ("mesh",) else path.suffix.lower().lstrip(".")
            empty_active = not any(f.obj == self.active_object for f in self.features)
            for i, part in enumerate(res.parts):
                reuse = i == 0 and empty_active and (len(res.parts) == 1 or len(self.objects) == 1)
                if reuse:
                    o = self.obj(self.active_object)
                    if len(res.parts) > 1 or part.name:
                        new = self._unique_obj_name(part.name) if part.name != o.name else o.name
                        if new != o.name:
                            self.set_object(o.name, checkpoint=False, new_name=new)
                            o = self.obj(new)
                    o.pose = [float(v) for v in part.pose]
                    if part.color:
                        o.color = part.color
                else:
                    o = self.add_object(self._unique_obj_name(part.name), part.color, checkpoint=False,
                                        activate=False, pose=part.pose)
                if part.fixed:
                    o.fixed = True
                fname = self._unique_feature_name(f"{part.name or path.stem}")
                ref = blob_of.get(part.proto) if part.proto else None
                if ref is None:
                    if getattr(part, "blob", None):                  # serialized by the import worker
                        ref = store_blob(part.blob)
                    else:
                        ref = store_blob(shape_to_text(part.shape))
                        fcache[ref[5:]] = part.shape          # rebuild uses it as is — no decode
                        register_geometry_key(part.shape, ref[5:])
                    if part.proto:
                        blob_of[part.proto] = ref
                self.add("import", {"data": ref, "source": path.name, "format": fmt},
                         name=fname, rebuild=False, checkpoint=False, obj=o.name)
                if progress and i % 50 == 0:
                    progress("Adding parts", i, len(res.parts))
                summary["objects"].append(o.name)
                summary["features"].append(fname)
            if res.parts:
                self.active_object = summary["objects"][0]
            if progress:
                progress("Building model", 0, 0)
            self.rebuild()
            if res.mates:
                added, notes = self._apply_imported_mates(res.mates, summary["objects"])
                summary["mates"] = added
                summary["notes"] += notes
        except Exception:
            if checkpoint:
                self.undo()
                self._redo.clear()
            raise
        self.dirty = True
        self._emit()
        return summary

    def _apply_imported_mates(self, data: dict, objects: list[str]) -> tuple[list[str], list[str]]:
        """Mates from the SolidWorks bridge. Entity geometry arrives in assembly (world)
        coordinates, mm; it is moved into each object's own frame here. Components are matched
        to imported objects by name, then by nearest bounding-box centre."""
        import numpy as np
        notes, added = [], []
        comps = data.get("components", [])
        cmap = self._match_components(comps, objects)
        for c in comps:
            o = cmap.get(c["name"])
            if o and c.get("fixed"):
                self.obj(o).fixed = True
        for m in data.get("mates", []):
            if m.get("type") not in asm.MATE_TYPES:
                notes.append(f"Mate {m.get('name', '?')} not imported: {m.get('type')} mates aren't supported")
                continue
            try:
                sides = []
                for side in (m["a"], m["b"]):
                    o = cmap.get(side["component"])
                    if o is None:
                        raise FeatureError(f"component '{side['component']}' not imported")
                    inv = np.linalg.inv(asm.pose_matrix(self.obj(o).pose))
                    ref = {"kind": side["kind"], "face": None}
                    ref["p"] = list((inv @ np.r_[np.asarray(side["p"], float), 1.0])[:3])
                    if side["kind"] == "plane":
                        ref["n"] = list(inv[:3, :3] @ np.asarray(side["n"], float))
                    elif side["kind"] == "axis":
                        ref["d"] = list(inv[:3, :3] @ np.asarray(side["d"], float))
                        ref["r"] = float(side.get("r", 0.0))
                    sides.append({"obj": o, "ref": ref})
                value = m.get("value")
                if value is None:
                    value = self._measure(m["type"], sides)
                mate = self.add_mate(m["type"], sides[0], sides[1], value=float(value or 0.0),
                                     flip=m.get("flip"), checkpoint=False)
                added.append(mate.name)
            except (FeatureError, KeyError, ValueError) as e:
                notes.append(f"Mate {m.get('name', '?')} ({m.get('type', '?')}) skipped: {e}")
        for s in data.get("skipped", []):
            notes.append(f"Mate {s.get('name', '?')} not imported: {s.get('reason', 'unsupported type')}")
        return added, notes

    def _measure(self, mtype, sides) -> float:
        """Current distance (mm) / angle (deg) between two refs as the parts sit now."""
        import numpy as np
        A = asm.world_ref(sides[0]["ref"], self.obj(sides[0]["obj"]).pose)
        B = asm.world_ref(sides[1]["ref"], self.obj(sides[1]["obj"]).pose)
        if mtype == "angle":
            u = np.asarray(A.get("n", A.get("d")), float)
            v = np.asarray(B.get("n", B.get("d")), float)
            return float(np.degrees(np.arccos(np.clip(abs(np.dot(u, v)), -1, 1))))
        d = np.asarray(B["p"], float) - np.asarray(A["p"], float)
        if A["kind"] == "plane":
            return float(abs(np.dot(d, A["n"])))
        return float(np.linalg.norm(d))

    def _match_components(self, comps: list[dict], objects: list[str]) -> dict:
        import numpy as np

        def norm(s):
            s = (s or "").split("/")[-1].strip().lower()
            return re.sub(r"[<>\s]", "", s)

        out, free = {}, list(objects)
        for c in comps:                                # exact names first
            o = next((o for o in free if norm(o) == norm(c["name"])), None)
            if o:
                out[c["name"]] = o
                free.remove(o)
        placed = self.placed_bodies()
        centres = {}
        for o in free:
            if o in placed:
                bb = placed[o].bounding_box()
                centres[o] = np.array([(bb.min.X + bb.max.X) / 2, (bb.min.Y + bb.max.Y) / 2,
                                       (bb.min.Z + bb.max.Z) / 2])
        for c in comps:
            if c["name"] in out or not centres or not c.get("box"):
                continue
            b = np.asarray(c["box"], float)
            cc = (b[:3] + b[3:]) / 2
            o = min(centres, key=lambda k: float(np.linalg.norm(centres[k] - cc)))
            if np.linalg.norm(centres[o] - cc) < 1.0 + 0.05 * float(np.linalg.norm(b[3:] - b[:3])):
                out[c["name"]] = o
                free.remove(o)
                del centres[o]
        for c in comps:                                # then names without the instance number
            if c["name"] in out:
                continue
            base = norm(c["name"]).rsplit("-", 1)[0]
            o = next((o for o in free if norm(o).rsplit("-", 1)[0] == base), None)
            if o:
                out[c["name"]] = o
                free.remove(o)
        return out


def _load_any(path: Path) -> "Document":
    """A .daedalus document, or any importable CAD file read into a fresh document."""
    path = Path(path)
    if not path.exists():
        raise FeatureError(f"{path} not found")
    if path.suffix.lower() in (".daedalus", ".strata"):
        return Document.open(path)
    d = Document()
    d.import_file(path)
    d.rebuild()
    return d
