"""v2.8 document features, mixed into Document:

* Global variables + equations that drive feature parameters and sketch dimensions.
* Configurations: named variants overriding parameters, suppression, variables, material and objects;
  a design table (CSV / Excel) in and out.
* Subassemblies (rigid or flexible groups of objects), component patterns (linear / circular) and
  mirrored components — instances share their source's geometry and follow it.
* Extra mate references (paths, cam profiles, slots) used by the v2.8 mates.

Qt-free.
"""
from __future__ import annotations

import ast
import copy
import math
import re
from pathlib import Path

import numpy as np

from . import assembly as asm
from .features import FEATURE_TYPES, FeatureError, coerce

DEFAULT_CONFIG = "Default"

# ------------------------------------------------------------------ expressions ---

_FUNCS = {
    "sin": lambda d: math.sin(math.radians(d)), "cos": lambda d: math.cos(math.radians(d)),
    "tan": lambda d: math.tan(math.radians(d)), "asin": lambda v: math.degrees(math.asin(v)),
    "acos": lambda v: math.degrees(math.acos(v)), "atan": lambda v: math.degrees(math.atan(v)),
    "atan2": lambda y, x: math.degrees(math.atan2(y, x)), "sqrt": math.sqrt, "abs": abs, "min": min, "max": max,
    "round": round, "floor": math.floor, "ceil": math.ceil, "int": int, "log": math.log, "exp": math.exp,
    "iif": lambda c, a, b: a if c else b, "if": lambda c, a, b: a if c else b,
}
_CONSTS = {"pi": math.pi, "e": math.e, "True": True, "False": False}


class ExprError(FeatureError):
    pass


def _dotted(node) -> str | None:
    if isinstance(node, ast.Name):
        return node.id
    if isinstance(node, ast.Attribute):
        base = _dotted(node.value)
        return f"{base}.{node.attr}" if base else None
    return None


def evaluate(expr, env: dict):
    """Safe arithmetic: numbers, + - * / // % **, comparisons, and/or/not, the functions in _FUNCS
    (trig in degrees), variables and dotted references (Extrude1.depth) looked up in env."""
    if isinstance(expr, (int, float)):
        return expr
    src = str(expr).strip().lstrip("=").replace("^", "**")
    src = re.sub(r'"([^"]+)"', lambda m: m.group(1).replace("@", ".").replace(" ", "_"), src)
    try:
        tree = ast.parse(src, mode="eval")
    except SyntaxError as e:
        raise ExprError(f"Can't read '{expr}': {e.msg}") from e

    def ev(n):
        if isinstance(n, ast.Expression):
            return ev(n.body)
        if isinstance(n, ast.Constant) and isinstance(n.value, (int, float, bool)):
            return n.value
        if isinstance(n, ast.BinOp):
            a, b = ev(n.left), ev(n.right)
            ops = {ast.Add: lambda: a + b, ast.Sub: lambda: a - b, ast.Mult: lambda: a * b,
                   ast.Div: lambda: a / b, ast.FloorDiv: lambda: a // b, ast.Mod: lambda: a % b,
                   ast.Pow: lambda: a ** b}
            f = ops.get(type(n.op))
            if f is None:
                raise ExprError("Operator not allowed")
            try:
                return f()
            except ZeroDivisionError as e:
                raise ExprError(f"Division by zero in '{expr}'") from e
        if isinstance(n, ast.UnaryOp):
            v = ev(n.operand)
            if isinstance(n.op, ast.USub):
                return -v
            if isinstance(n.op, ast.UAdd):
                return +v
            if isinstance(n.op, ast.Not):
                return not v
        if isinstance(n, ast.BoolOp):
            vals = [ev(x) for x in n.values]
            return all(vals) if isinstance(n.op, ast.And) else any(vals)
        if isinstance(n, ast.Compare):
            left = ev(n.left)
            for op, c in zip(n.ops, n.comparators):
                right = ev(c)
                ok = {ast.Lt: left < right, ast.LtE: left <= right, ast.Gt: left > right, ast.GtE: left >= right,
                      ast.Eq: left == right, ast.NotEq: left != right}.get(type(op))
                if ok is None:
                    raise ExprError("Comparison not allowed")
                if not ok:
                    return False
                left = right
            return True
        if isinstance(n, ast.IfExp):
            return ev(n.body) if ev(n.test) else ev(n.orelse)
        if isinstance(n, ast.Call):
            fname = n.func.id if isinstance(n.func, ast.Name) else None
            if fname not in _FUNCS:
                raise ExprError(f"Unknown function '{fname}'")
            return _FUNCS[fname](*[ev(a) for a in n.args])
        if isinstance(n, (ast.Name, ast.Attribute)):
            key = _dotted(n)
            if key in env:
                return env[key]
            if key in _CONSTS:
                return _CONSTS[key]
            raise ExprError(f"Unknown name '{key}' in '{expr}'")
        raise ExprError(f"Not allowed in an equation: {type(n).__name__}")

    v = ev(tree)
    if isinstance(v, bool):
        return v
    if not isinstance(v, (int, float)) or (isinstance(v, float) and not math.isfinite(v)):
        raise ExprError(f"'{expr}' doesn't give a number")
    return v


def names_in(expr) -> set:
    src = str(expr).strip().lstrip("=").replace("^", "**")
    src = re.sub(r'"([^"]+)"', lambda m: m.group(1).replace("@", ".").replace(" ", "_"), src)
    try:
        tree = ast.parse(src, mode="eval")
    except SyntaxError:
        return set()
    out = set()
    for n in ast.walk(tree):
        if isinstance(n, (ast.Name, ast.Attribute)):
            k = _dotted(n)
            if k and k not in _FUNCS and k not in _CONSTS:
                out.add(k)
    return {k for k in out if not any(o != k and o.startswith(k + ".") for o in out)}


# ------------------------------------------------------------- target access ---

_DIM = re.compile(r"^D(\d+)$")


def _split_target(target: str):
    if "." not in target:
        raise FeatureError(f"Target '{target}' must be Feature.param (e.g. Extrude1.depth or Sketch1.D2)")
    feat, key = target.split(".", 1)
    return feat, key


def _get_value(f, key):
    m = _DIM.match(key)
    if m and f.TYPE in ("sketch", "sketch3d"):
        cons = (f.params.get("geometry") or {}).get("constraints", [])
        i = int(m.group(1))
        if not 0 <= i < len(cons) or "value" not in cons[i]:
            raise FeatureError(f"{f.name} has no dimension D{i}")
        return cons[i]["value"]
    if f.TYPE in ("sketch", "sketch3d"):
        for c in (f.params.get("geometry") or {}).get("constraints", []):
            if c.get("name") == key and "value" in c:
                return c["value"]
    if key not in f.params:
        raise FeatureError(f"{f.name} has no parameter '{key}'")
    return f.params[key]


def _set_value(f, key, value):
    m = _DIM.match(key)
    named = None
    if f.TYPE in ("sketch", "sketch3d") and not m:
        cons = (f.params.get("geometry") or {}).get("constraints", [])
        named = next((i for i, c in enumerate(cons) if c.get("name") == key and "value" in c), None)
    if (m or named is not None) and f.TYPE in ("sketch", "sketch3d"):
        geom = copy.deepcopy(f.params.get("geometry") or {})
        cons = geom.get("constraints", [])
        i = int(m.group(1)) if m else named
        if not 0 <= i < len(cons) or "value" not in cons[i]:
            raise FeatureError(f"{f.name} has no dimension D{i}")
        cons[i]["value"] = float(value)
        f.params["geometry"] = geom
        return
    spec = f.param_spec(key)
    f.params[key] = coerce(spec, value)


def _numeric_params(f) -> dict:
    out = {}
    for p in f.PARAMS:
        v = f.params.get(p.name)
        if p.kind in ("float", "int", "bool") and isinstance(v, (int, float, bool)):
            out[f"{f.name}.{p.name}"] = v
    if f.TYPE in ("sketch", "sketch3d"):
        for i, c in enumerate((f.params.get("geometry") or {}).get("constraints", [])):
            if "value" in c:
                out[f"{f.name}.D{i}"] = c["value"]
                if c.get("name"):
                    out[f"{f.name}.{c['name']}"] = c["value"]
    return out


def _clone(f):
    g = FEATURE_TYPES[f.TYPE].__new__(FEATURE_TYPES[f.TYPE])
    g.name, g.suppressed, g.layer, g.obj = f.name, f.suppressed, f.layer, f.obj
    g.params = copy.deepcopy(f.params)
    g._origin = f
    return g


# --------------------------------------------------------------------- mixin ---

class V28Mixin:
    # ------------------------------------------------------------ state
    def _v28_init(self):
        self.variables: dict = {}             # name -> expression
        self.equations: list = []             # [{"target", "expr", "suppressed"}]
        self.configurations: list = []        # [{"name", "description", "params", "suppress", "variables", ...}]
        self.active_config = DEFAULT_CONFIG
        self.subassemblies: list = []         # [{"name", "parent", "rigid"}]
        self.comp_patterns: list = []         # [{"name", "kind", "seeds", ..., "instances"}]
        self.eq_values: dict = {}             # last evaluated equation values (target -> value)
        self.eq_errors: dict = {}
        from . import units as _U
        self.units: dict = dict(type(self).default_units or {"name": _U.DEFAULT})
        _U.set_active(self.units)

    def _v28_dict(self) -> dict:
        return {"variables": dict(self.variables), "equations": [dict(e) for e in self.equations],
                "configurations": copy.deepcopy(self.configurations), "active_config": self.active_config,
                "subassemblies": [dict(s) for s in self.subassemblies],
                "comp_patterns": copy.deepcopy(self.comp_patterns), "units": dict(self.units)}

    def _v28_load(self, d: dict):
        self._v28_init()
        self.variables = dict(d.get("variables", {}))
        self.equations = [dict(e) for e in d.get("equations", [])]
        self.configurations = copy.deepcopy(d.get("configurations", []))
        self.active_config = d.get("active_config", DEFAULT_CONFIG)
        if self.active_config != DEFAULT_CONFIG and not any(c["name"] == self.active_config
                                                            for c in self.configurations):
            self.active_config = DEFAULT_CONFIG
        self.subassemblies = [dict(s) for s in d.get("subassemblies", [])]
        self.comp_patterns = copy.deepcopy(d.get("comp_patterns", []))
        from . import units as _U
        self.units = dict(d.get("units") or {"name": _U.DEFAULT})
        _U.set_active(self.units)

    # ========================================================== units
    default_units: dict | None = None          # new documents (the app sets this from its settings)

    def unit_system(self):
        from . import units as _U
        return _U.system_from(self.units)

    def set_units(self, name: str | None = None, units: dict | None = None, decimals: dict | None = None,
                  length_format: str | None = None, denominator: int | None = None, reset: bool = False):
        """Document unit system (display + typed input; geometry is always stored in mm/deg).
        name: a preset in units.SYSTEMS; units {quantity: unit} overrides; decimals {length|angle|default|<q>:
        n}; length_format decimal|fractional|ft-in; denominator 2…128 for fractions."""
        from . import units as _U
        cur = {} if reset else dict(self.units)
        if name is not None:
            key = next((k for k, v in _U.SYSTEMS.items()
                        if name in (k, v.label) or name.lower() in (k.lower(), v.label.lower())), None)
            if key is None:
                raise FeatureError(f"Unknown unit system '{name}'. Systems: {list(_U.SYSTEMS)}")
            cur = {"name": key}
        cur.setdefault("name", _U.DEFAULT)
        if units:
            ov = dict(cur.get("units") or {})
            for q, u in units.items():
                if q not in _U.Q:
                    raise FeatureError(f"Unknown quantity '{q}'. Quantities: {list(_U.Q)}")
                try:
                    ov[q] = _U.canon(u, q)
                except _U.UnitError as e:
                    raise FeatureError(f"{e}. Units: {_U.units_of(q)}") from e
            cur["units"] = ov
        if decimals:
            cur["decimals"] = {**(cur.get("decimals") or {}), **{k: max(0, min(8, int(v))) for k, v in
                                                                  decimals.items()}}
        if length_format is not None:
            if length_format not in ("decimal", "fractional", "ft-in"):
                raise FeatureError("length_format: decimal | fractional | ft-in")
            cur["length_format"] = length_format
        if denominator is not None:
            if int(denominator) not in (2, 4, 8, 16, 32, 64, 128):
                raise FeatureError("denominator: 2, 4, 8, 16, 32, 64 or 128")
            cur["denominator"] = int(denominator)
        self.units = cur
        _U.set_active(cur)
        self.dirty = True
        self._emit()
        return _U.system_from(cur)

    # ======================================================= equations
    def config(self, name: str | None = None) -> dict:
        name = name or self.active_config
        if name == DEFAULT_CONFIG:
            return {"name": DEFAULT_CONFIG, "params": {}, "suppress": {}, "variables": {}, "objects": {}}
        for c in self.configurations:
            if c["name"] == name:
                return c
        raise FeatureError(f"No configuration '{name}'. Configurations: {self.config_names()}")

    def config_names(self) -> list:
        return [DEFAULT_CONFIG] + [c["name"] for c in self.configurations]

    def _env_variables(self, cfg: dict, env: dict, errors: dict):
        """Resolve variables (in dependency order) into env."""
        exprs = dict(self.variables)
        exprs.update(cfg.get("variables", {}))
        pending = dict(exprs)
        for _ in range(len(pending) + 1):
            if not pending:
                break
            progressed = False
            for k in list(pending):
                need = names_in(pending[k]) - set(env) - set(_CONSTS)
                if need & set(pending):
                    continue
                try:
                    env[k] = evaluate(pending[k], env)
                except FeatureError as e:
                    errors[k] = str(e)
                pending.pop(k)
                progressed = True
            if not progressed:
                for k in pending:
                    errors[k] = "Circular reference between variables"
                break

    def set_variable(self, name: str, expr, checkpoint=True):
        if not re.fullmatch(r"[A-Za-z_]\w*", name or ""):
            raise FeatureError("Variable names: letters, digits and _ (not starting with a digit)")
        if name in _FUNCS or name in _CONSTS:
            raise FeatureError(f"'{name}' is reserved")
        env = {}
        self._env_variables({"variables": {}}, env, {})
        env.update({k: 1.0 for k in self.variables})
        if checkpoint:
            self.checkpoint()
        self.variables[name] = str(expr)
        self.dirty = True
        self.rebuild()

    def delete_variable(self, name: str, checkpoint=True):
        if name not in self.variables:
            raise FeatureError(f"No variable '{name}'")
        if checkpoint:
            self.checkpoint()
        self.variables.pop(name)
        self.dirty = True
        self.rebuild()

    def add_equation(self, target: str, expr: str, checkpoint=True, index: int | None = None) -> dict:
        feat, key = _split_target(target)
        f = self.get(feat)
        _get_value(f, key)
        names_in(expr)
        evaluate_check = expr                                         # syntax check
        try:
            ast.parse(str(evaluate_check).lstrip("=").replace("^", "**"), mode="eval")
        except SyntaxError as e:
            raise FeatureError(f"Can't read '{expr}': {e.msg}") from e
        if any(e["target"] == target for e in self.equations):
            raise FeatureError(f"{target} is already driven by an equation")
        if checkpoint:
            self.checkpoint()
        eq = {"target": target, "expr": str(expr), "suppressed": False}
        if index is None:
            self.equations.append(eq)
        else:
            self.equations.insert(index, eq)
        self.dirty = True
        self.rebuild()
        if target in self.eq_errors:
            err = self.eq_errors[target]
            self.equations.remove(eq)
            if checkpoint and self._undo:
                self._undo.pop()
            self.rebuild()
            raise FeatureError(f"Equation for {target}: {err}")
        return eq

    def set_equation(self, target: str, expr: str | None = None, suppressed: bool | None = None, checkpoint=True):
        eq = next((e for e in self.equations if e["target"] == target), None)
        if eq is None:
            raise FeatureError(f"No equation drives {target}")
        if checkpoint:
            self.checkpoint()
        if expr is not None:
            eq["expr"] = str(expr)
        if suppressed is not None:
            eq["suppressed"] = bool(suppressed)
        self.dirty = True
        self.rebuild()

    def delete_equation(self, target: str, checkpoint=True):
        eq = next((e for e in self.equations if e["target"] == target), None)
        if eq is None:
            raise FeatureError(f"No equation drives {target}")
        if checkpoint:
            self.checkpoint()
        self.equations.remove(eq)
        self.dirty = True
        self.rebuild()

    def driven_params(self) -> dict:
        """{feature: {param: expression}} for equation-driven values (Inspector greys them out)."""
        out = {}
        for e in self.equations:
            if e.get("suppressed"):
                continue
            f, k = e["target"].split(".", 1)
            out.setdefault(f, {})[k] = e["expr"]
        return out

    def effective_features(self) -> list:
        """The history as the active configuration + equations make it. Without either, the features
        themselves (no copies)."""
        self.eq_values, self.eq_errors = {}, {}
        cfg = self.config() if self.active_config != DEFAULT_CONFIG else None
        eqs = [e for e in self.equations if not e.get("suppressed")]
        if cfg is None and not eqs and not self.variables:
            return self.features
        if cfg is None and not eqs:
            return self.features
        cfg = cfg or self.config(DEFAULT_CONFIG)
        out = [_clone(f) for f in self.features]
        by = {f.name: f for f in out}
        hidden_objs = {o for o, st in cfg.get("objects", {}).items() if st.get("suppressed")}
        for f in out:
            if f.name in cfg.get("suppress", {}):
                f.suppressed = bool(cfg["suppress"][f.name])
            if f.obj in hidden_objs:
                f.suppressed = True
        for target, v in cfg.get("params", {}).items():
            try:
                fn, key = _split_target(target)
                if fn in by:
                    _set_value(by[fn], key, v)
            except FeatureError as e:
                self.eq_errors[target] = str(e)
        env = {}
        for f in out:
            env.update(_numeric_params(f))
        self._env_variables(cfg, env, self.eq_errors)
        for e in eqs:
            try:
                fn, key = _split_target(e["target"])
                if fn not in by:
                    raise FeatureError(f"No feature '{fn}'")
                val = evaluate(e["expr"], env)
                _set_value(by[fn], key, val)
                self.eq_values[e["target"]] = val
                env.update(_numeric_params(by[fn]))
            except FeatureError as ex:
                self.eq_errors[e["target"]] = str(ex)
        return out

    def variable_values(self) -> dict:
        env, err = {}, {}
        for f in self.features:
            env.update(_numeric_params(f))
        cfg = self.config() if self.active_config != DEFAULT_CONFIG else {"variables": {}}
        self._env_variables(cfg, env, err)
        return {k: env.get(k) for k in {**self.variables, **cfg.get("variables", {})}}

    # =================================================== configurations
    def add_configuration(self, name: str, description: str = "", copy_from: str | None = None,
                          activate=True, checkpoint=True) -> dict:
        if not re.fullmatch(r"[\w\- ]+", name or "") or name in self.config_names():
            raise FeatureError(f"Configuration name '{name}' is invalid or taken")
        if checkpoint:
            self.checkpoint()
        base = copy.deepcopy(self.config(copy_from)) if copy_from else {}
        c = {"name": name, "description": description, "params": base.get("params", {}),
             "suppress": base.get("suppress", {}), "variables": base.get("variables", {}),
             "material": base.get("material"), "objects": base.get("objects", {}),
             "props": base.get("props", {})}
        self.configurations.append(c)
        if activate:
            self.active_config = name
        self.dirty = True
        self.rebuild()
        return c

    def set_configuration(self, name: str, checkpoint=True, params: dict | None = None,
                          suppress: dict | None = None, variables: dict | None = None, material=None,
                          objects: dict | None = None, description: str | None = None, new_name: str | None = None,
                          props: dict | None = None):
        """Merge overrides into a configuration (value None removes that override)."""
        if name == DEFAULT_CONFIG:
            raise FeatureError("The Default configuration has no overrides — edit the model itself")
        c = self.config(name)
        if checkpoint:
            self.checkpoint()
        for key, upd in (("params", params), ("suppress", suppress), ("variables", variables),
                         ("objects", objects), ("props", props)):
            if upd:
                for k, v in upd.items():
                    if v is None:
                        c.setdefault(key, {}).pop(k, None)
                    else:
                        if key == "params":
                            fn, pk = _split_target(k)
                            _get_value(self.get(fn), pk)
                        if key == "suppress":
                            self.get(k)
                        c.setdefault(key, {})[k] = v
        if material is not None:
            c["material"] = material or None
        if description is not None:
            c["description"] = description
        if new_name and new_name != name:
            if new_name in self.config_names():
                raise FeatureError(f"'{new_name}' is taken")
            c["name"] = new_name
            if self.active_config == name:
                self.active_config = new_name
        self.dirty = True
        self.rebuild()
        return c

    def delete_configuration(self, name: str, checkpoint=True):
        c = self.config(name)
        if name == DEFAULT_CONFIG:
            raise FeatureError("Default can't be deleted")
        if checkpoint:
            self.checkpoint()
        self.configurations.remove(c)
        if self.active_config == name:
            self.active_config = DEFAULT_CONFIG
        self.dirty = True
        self.rebuild()

    def activate_configuration(self, name: str, checkpoint=False):
        self.config(name)
        if checkpoint:
            self.checkpoint()
        self.active_config = name
        self.dirty = True
        self.rebuild()

    def config_material(self, obj: str | None = None):
        if self.active_config == DEFAULT_CONFIG:
            return None
        c = self.config()
        if obj and obj in c.get("objects", {}) and c["objects"][obj].get("material"):
            return c["objects"][obj]["material"]
        return c.get("material")

    # ---------------------------------------------------- design table
    def config_table(self) -> tuple[list, list]:
        """(columns, rows) — one row per configuration. Columns: Configuration, Description, then
        $VAR:<name>, <Feature.param>, $STATE@<Feature> (S = suppressed, U = unsuppressed), $MATERIAL."""
        cols = ["Configuration", "Description"]
        keys = []
        for c in self.configurations:
            for k in c.get("variables", {}):
                keys.append(f"$VAR:{k}")
            for k in c.get("params", {}):
                keys.append(k)
            for k in c.get("suppress", {}):
                keys.append(f"$STATE@{k}")
            if c.get("material"):
                keys.append("$MATERIAL")
        for k in keys:
            if k not in cols:
                cols.append(k)
        rows = []
        for name in self.config_names():
            c = self.config(name)
            row = {"Configuration": name, "Description": c.get("description", "")}
            for k in cols[2:]:
                if k.startswith("$VAR:"):
                    row[k] = c.get("variables", {}).get(k[5:], self.variables.get(k[5:], ""))
                elif k.startswith("$STATE@"):
                    fn = k[7:]
                    sup = c.get("suppress", {}).get(fn)
                    if sup is None:
                        sup = self.get(fn).suppressed if fn in self._fnames() else False
                    row[k] = "S" if sup else "U"
                elif k == "$MATERIAL":
                    row[k] = c.get("material") or self.material
                else:
                    fn, pk = _split_target(k)
                    v = c.get("params", {}).get(k)
                    row[k] = v if v is not None else (_get_value(self.get(fn), pk) if fn in self._fnames() else "")
            rows.append(row)
        return cols, rows

    def import_config_table(self, rows: list[dict], checkpoint=True) -> list[str]:
        """Create / update configurations from design-table rows. Values equal to the model's own are not
        stored as overrides. Returns the configurations touched."""
        if checkpoint:
            self.checkpoint()
        touched = []
        for row in rows:
            name = str(row.get("Configuration") or "").strip()
            if not name or name == DEFAULT_CONFIG:
                continue
            if name not in self.config_names():
                self.configurations.append({"name": name, "description": "", "params": {}, "suppress": {},
                                            "variables": {}, "material": None, "objects": {}, "props": {}})
            c = self.config(name)
            c["description"] = str(row.get("Description") or c.get("description", ""))
            for k, v in row.items():
                if k in ("Configuration", "Description") or v in (None, ""):
                    continue
                if k.startswith("$VAR:"):
                    c["variables"][k[5:]] = str(v)
                elif k.startswith("$STATE@"):
                    c["suppress"][k[7:]] = str(v).strip().upper().startswith("S")
                elif k == "$MATERIAL":
                    c["material"] = str(v) if str(v) != self.material else None
                else:
                    fn, pk = _split_target(k)
                    f = self.get(fn)
                    base = _get_value(f, pk)
                    try:
                        val = float(v) if isinstance(base, (int, float)) and not isinstance(base, bool) else v
                    except ValueError as e:
                        raise FeatureError(f"{name}: {k} = {v!r} isn't a number") from e
                    if val != base:
                        c["params"][k] = val
                    else:
                        c["params"].pop(k, None)
            touched.append(name)
        self.dirty = True
        self.rebuild()
        return touched

    def export_config_table(self, path):
        cols, rows = self.config_table()
        path = Path(path)
        if path.suffix.lower() == ".xlsx":
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "Design Table"
            ws.append(cols)
            for r in rows:
                ws.append([r.get(c, "") for c in cols])
            wb.save(path)
        else:
            import csv
            with open(path, "w", newline="", encoding="utf-8") as fh:
                w = csv.DictWriter(fh, fieldnames=cols)
                w.writeheader()
                w.writerows(rows)
        return str(path)

    def read_config_table(self, path, checkpoint=True):
        path = Path(path)
        if path.suffix.lower() == ".xlsx":
            from openpyxl import load_workbook
            ws = load_workbook(path, data_only=True).active
            it = ws.iter_rows(values_only=True)
            cols = [str(c) if c is not None else "" for c in next(it)]
            rows = [dict(zip(cols, r)) for r in it if any(v not in (None, "") for v in r)]
        else:
            import csv
            with open(path, newline="", encoding="utf-8-sig") as fh:
                rows = list(csv.DictReader(fh))
        return self.import_config_table(rows, checkpoint=checkpoint)

    # ==================================================== subassemblies
    def subassembly(self, name: str) -> dict:
        for s in self.subassemblies:
            if s["name"] == name:
                return s
        raise FeatureError(f"No subassembly '{name}'")

    def subassembly_members(self, name: str, recursive=True) -> list[str]:
        names = {name}
        if recursive:
            changed = True
            while changed:
                changed = False
                for s in self.subassemblies:
                    if s.get("parent") in names and s["name"] not in names:
                        names.add(s["name"])
                        changed = True
        return [o.name for o in self.objects if o.group in names]

    def make_subassembly(self, name: str | None, objects: list[str], rigid=True, parent: str | None = None,
                         checkpoint=True) -> dict:
        if not objects:
            raise FeatureError("Pick the objects for the subassembly")
        for o in objects:
            self.obj(o)
        if not name:
            n = 1
            while any(s["name"] == f"Subassembly{n}" for s in self.subassemblies):
                n += 1
            name = f"Subassembly{n}"
        if any(s["name"] == name for s in self.subassemblies) or self.has_obj(name):
            raise FeatureError(f"Name '{name}' is taken")
        if parent:
            self.subassembly(parent)
        if checkpoint:
            self.checkpoint()
        s = {"name": name, "parent": parent, "rigid": bool(rigid)}
        # objects already in a subassembly: that subassembly becomes the parent (nesting)
        groups = {self.obj(o).group for o in objects} - {None}
        if not parent and len(groups) == 1:
            s["parent"] = next(iter(groups))
        self.subassemblies.append(s)
        for o in objects:
            self.obj(o).group = name
        self.dirty = True
        self.solve_assembly()
        self._emit()
        return s

    def set_subassembly(self, name: str, checkpoint=True, rigid: bool | None = None, new_name: str | None = None,
                        parent=...):
        s = self.subassembly(name)
        if checkpoint:
            self.checkpoint()
        if rigid is not None:
            s["rigid"] = bool(rigid)
        if parent is not ...:
            if parent == name or (parent and name in self._sub_ancestors(parent)):
                raise FeatureError("A subassembly can't contain itself")
            s["parent"] = parent or None
        if new_name and new_name != name:
            if any(x["name"] == new_name for x in self.subassemblies):
                raise FeatureError(f"'{new_name}' is taken")
            for o in self.objects:
                if o.group == name:
                    o.group = new_name
            for x in self.subassemblies:
                if x.get("parent") == name:
                    x["parent"] = new_name
            s["name"] = new_name
        self.dirty = True
        self.solve_assembly()
        self._emit()

    def _sub_ancestors(self, name):
        out, cur = [], name
        while cur:
            out.append(cur)
            cur = next((s.get("parent") for s in self.subassemblies if s["name"] == cur), None)
        return out

    def dissolve_subassembly(self, name: str, checkpoint=True):
        s = self.subassembly(name)
        if checkpoint:
            self.checkpoint()
        for o in self.objects:
            if o.group == name:
                o.group = s.get("parent")
        for x in self.subassemblies:
            if x.get("parent") == name:
                x["parent"] = s.get("parent")
        self.subassemblies.remove(s)
        self.dirty = True
        self.solve_assembly()
        self._emit()

    def move_to_subassembly(self, objects: list[str], name: str | None, checkpoint=True):
        if name:
            self.subassembly(name)
        if checkpoint:
            self.checkpoint()
        for o in objects:
            self.obj(o).group = name or None
        self.dirty = True
        self.solve_assembly()
        self._emit()

    def _rigid_groups(self) -> list[list[str]]:
        """Objects that move as one: each rigid subassembly with everything nested in it."""
        out = []
        for s in self.subassemblies:
            if not s.get("rigid", True):
                continue
            if any(p != s["name"] and self.subassembly(p).get("rigid", True)
                   for p in self._sub_ancestors(s["name"])[1:] if any(x["name"] == p for x in self.subassemblies)):
                continue                                     # a rigid parent already covers it
            mem = [n for n in self.subassembly_members(s["name"]) if not self.obj(n).suppressed]
            if len(mem) > 1:
                out.append(mem)
        return out

    def save_subassembly(self, name: str, path) -> str:
        """Write a subassembly (its objects, their features and the mates between them) as a .daedalus file."""
        import json
        mem = set(self.subassembly_members(name))
        d = self.to_dict()
        d["objects"] = [o for o in d["objects"] if o["name"] in mem]
        for o in d["objects"]:
            o["group"] = None
        d["active_object"] = d["objects"][0]["name"]
        d["features"] = [f for f in d["features"] if f.get("obj") in mem]
        d["mates"] = [m for m in d["mates"] if {m["a"]["obj"], m["b"]["obj"]} <= mem]
        d["subassemblies"] = [s for s in d.get("subassemblies", []) if s["name"] != name and
                              any(self.obj(o).group == s["name"] for o in mem)]
        d["comp_patterns"] = [p for p in d.get("comp_patterns", []) if set(p["seeds"]) <= mem]
        d["drawings"], d["studies"], d["configurations"], d["equations"] = [], [], [], []
        Path(path).write_text(json.dumps(d, indent=1))
        return str(path)

    def tree(self) -> list:
        """Assembly structure: [{"name", "kind": "subassembly"|"object", "children": [...]}] top-down."""
        def node(sname):
            kids = [node(s["name"]) for s in self.subassemblies if s.get("parent") == sname]
            kids += [{"name": o.name, "kind": "object"} for o in self.objects if o.group == sname]
            return {"name": sname, "kind": "subassembly", "rigid": self.subassembly(sname).get("rigid", True),
                    "children": kids}
        out = [node(s["name"]) for s in self.subassemblies if not s.get("parent")]
        out += [{"name": o.name, "kind": "object"} for o in self.objects if not o.group]
        return out

    # =============================================== component patterns
    def comp_pattern(self, name: str) -> dict:
        for p in self.comp_patterns:
            if p["name"] == name:
                return p
        raise FeatureError(f"No component pattern '{name}'")

    def add_component_pattern(self, kind: str, seeds: list[str], checkpoint=True, name: str | None = None,
                              **params) -> dict:
        """kind linear: dir1 [x,y,z], spacing1, count1, dir2, spacing2, count2 ·
        circular: axis {"p": [x,y,z], "d": [x,y,z]} (or "X"/"Y"/"Z"), count, angle (360 = equal spacing) ·
        mirror: plane {"p": [x,y,z], "n": [x,y,z]} (or "XY"/"XZ"/"YZ")."""
        if kind not in ("linear", "circular", "mirror"):
            raise FeatureError("kind: linear | circular | mirror")
        if not seeds:
            raise FeatureError("Pick the components to pattern")
        for s_ in seeds:
            o = self.obj(s_)
            if o.pattern:
                raise FeatureError(f"{s_} is itself a pattern instance — pattern its seed instead")
        if not name:
            n = 1
            base = {"linear": "LocalLPattern", "circular": "LocalCirPattern", "mirror": "MirrorComponent"}[kind]
            while any(p["name"] == f"{base}{n}" for p in self.comp_patterns):
                n += 1
            name = f"{base}{n}"
        pat = {"name": name, "kind": kind, "seeds": list(seeds), "instances": []}
        pat.update(_pattern_defaults(kind))
        pat.update({k: v for k, v in params.items() if v is not None})
        _check_pattern(pat)
        if checkpoint:
            self.checkpoint()
        self.comp_patterns.append(pat)
        self._regen_pattern(pat)
        self.dirty = True
        self.rebuild()
        return pat

    def set_component_pattern(self, name: str, checkpoint=True, **params):
        pat = self.comp_pattern(name)
        new = dict(pat)
        new.update({k: v for k, v in params.items() if v is not None})
        _check_pattern(new)
        if checkpoint:
            self.checkpoint()
        pat.update(new)
        self._regen_pattern(pat)
        self.dirty = True
        self.rebuild()

    def delete_component_pattern(self, name: str, checkpoint=True):
        pat = self.comp_pattern(name)
        if checkpoint:
            self.checkpoint()
        self._drop_instances(pat)
        self.comp_patterns.remove(pat)
        self.dirty = True
        self.rebuild()

    def _drop_instances(self, pat):
        gone = set(pat.get("instances", []))
        self.objects = [o for o in self.objects if o.name not in gone]
        self.mates = [m for m in self.mates if not set(m.objects()) & gone]
        self.physics.drives = [d for d in self.physics.drives if any(m.name == d.mate for m in self.mates)]
        if self.active_object in gone:
            self.active_object = self.objects[0].name
        self._names_changed()
        pat["instances"] = []

    def _pattern_transforms(self, pat) -> list:
        """World transforms (4×4) of each instance slot, excluding the seed itself."""
        out = []
        if pat["kind"] == "linear":
            d1 = _unit(_axis_vec(pat["dir1"]))
            d2 = _unit(_axis_vec(pat["dir2"]))
            for i in range(int(pat["count1"])):
                for j in range(int(pat["count2"])):
                    if i == 0 and j == 0:
                        continue
                    M = np.eye(4)
                    M[:3, 3] = d1 * float(pat["spacing1"]) * i + d2 * float(pat["spacing2"]) * j
                    out.append(M)
        elif pat["kind"] == "circular":
            p, d = _axis_spec(pat["axis"])
            n = int(pat["count"])
            ang = float(pat["angle"])
            step = ang / n if abs(ang - 360) < 1e-9 else ang / max(n - 1, 1)
            for i in range(1, n):
                R = asm.rotmat(d * math.radians(step * i))
                M = np.eye(4)
                M[:3, :3] = R
                M[:3, 3] = p - R @ p
                out.append(M)
        else:
            out.append(None)                                 # one mirrored copy (special-cased)
        return out

    def _regen_pattern(self, pat):
        self._drop_instances(pat)
        slots = self._pattern_transforms(pat)
        made = []
        for seed in pat["seeds"]:
            so = self.obj(seed)
            for k, _M in enumerate(slots):
                base = f"{seed}-{'M' if pat['kind'] == 'mirror' else ''}{k + 1}" if pat["kind"] != "mirror" \
                    else f"{seed}-Mirror"
                nm = base
                n = 2
                while self.has_obj(nm):
                    nm = f"{base}_{n}"
                    n += 1
                o = asm.Obj(nm, so.color, so.visible, True, list(so.pose), None, so.part_no, so.unit_cost,
                            so.description, copy.deepcopy(so.std), None, so.group,
                            instance_of=seed, mirrored=pat["kind"] == "mirror", pattern=pat["name"])
                self.objects.append(o)
                self._noted_add("obj", o)
                made.append(nm)
        pat["instances"] = made
        self._update_pattern_poses()

    def _update_pattern_poses(self):
        """Instances follow their seeds (call after every assembly solve)."""
        for pat in self.comp_patterns:
            slots = self._pattern_transforms(pat)
            inst = iter(pat.get("instances", []))
            for seed in pat["seeds"]:
                if not self.has_obj(seed):
                    continue
                Ms = asm.pose_matrix(self.obj(seed).pose)
                for M in slots:
                    nm = next(inst, None)
                    if nm is None or not self.has_obj(nm):
                        continue
                    o = self.obj(nm)
                    if pat["kind"] == "mirror":
                        p, n = _plane_spec(pat["plane"])
                        H = np.eye(3) - 2 * np.outer(n, n)
                        S = np.diag([1.0, 1.0, -1.0])
                        R = H @ Ms[:3, :3] @ S
                        t = H @ Ms[:3, 3] + 2 * n * float(n @ p)
                        o.pose = [float(v) for v in t] + [float(v) for v in asm.rotvec(R)]
                    else:
                        W = M @ Ms
                        o.pose = [float(v) for v in W[:3, 3]] + [float(v) for v in asm.rotvec(W[:3, :3])]

    def instance_body(self, o, bodies: dict):
        src = bodies.get(o.instance_of)
        if src is None:
            return None
        if not o.mirrored:
            return src
        cache = self.__dict__.setdefault("_mirror_cache", {})
        hit = cache.get(o.instance_of)
        if hit is not None and hit[0] is src:
            return hit[1]
        from build123d import Plane, mirror
        mb = mirror(src, about=Plane.XY)
        cache[o.instance_of] = (src, mb)
        return mb

    # ============================================ special mate references
    def _ref_special(self, obj: str, spec):
        """path / cam profile / slot references for the v2.8 mates, or None."""
        if not isinstance(spec, dict) or "kind" in spec:
            return None
        body = self.result.bodies.get(obj)
        if "path" in spec:                                   # a sketch / 3D sketch / curve feature's path
            import numpy as _np
            src = spec["path"]
            ctx = self._ctx_of(obj)
            wires = ctx.paths.get(src) if ctx else None
            if not wires:
                raise FeatureError(f"'{src}' has no path curves on {obj}")
            w = max(wires, key=lambda w_: w_.length)
            n = max(16, min(400, int(w.length / 1.0)))
            pts = [w.position_at(i / n) for i in range(n + 1)]
            P = _np.array([[q.X, q.Y, q.Z] for q in pts])
            closed = bool(getattr(w, "is_closed", False))
            return {"kind": "path", "p": list(P[0]), "pts": P.tolist(), "closed": closed, "face": None,
                    "source": src}
        if body is None:
            return None
        if "edge_near" in spec:
            p = tuple(float(v) for v in spec["edge_near"])
            e = min(body.edges(), key=lambda e_: e_.distance_to(p))
            n = 2 if e.geom_type.name == "LINE" else max(16, min(200, int(e.length)))
            pts = [e.position_at(i / n) for i in range(n + 1)]
            P = [[q.X, q.Y, q.Z] for q in pts]
            return {"kind": "path", "p": P[0], "pts": P, "closed": bool(e.is_closed), "face": None}
        if "cam_face" in spec:
            return cam_profile(body, int(spec["cam_face"]), spec.get("pick"))
        if "slot_face" in spec:
            return slot_ref(body, int(spec["slot_face"]))
        return None

    def _ctx_of(self, obj):
        """Build context at the end of an object's history (for its sketch paths)."""
        from .features import BuildContext
        ctx = BuildContext()
        off = {l.name for l in self.layers if not l.in_model}
        for g in self.effective_features():
            if g.obj != obj or g.suppressed or g.layer in off:
                continue
            before = ctx.body
            try:
                g.build(ctx)
            except Exception:  # noqa: BLE001
                ctx.body = before
        return ctx


# --------------------------------------------------------------- helpers ---

def _pattern_defaults(kind):
    if kind == "linear":
        return {"dir1": [1, 0, 0], "spacing1": 50.0, "count1": 3, "dir2": [0, 1, 0], "spacing2": 50.0, "count2": 1}
    if kind == "circular":
        return {"axis": "Z", "count": 4, "angle": 360.0}
    return {"plane": "YZ"}


def _check_pattern(p):
    if p["kind"] == "linear":
        if int(p["count1"]) < 1 or int(p["count2"]) < 1 or int(p["count1"]) * int(p["count2"]) < 2:
            raise FeatureError("A linear pattern needs at least 2 instances")
        if int(p["count1"]) * int(p["count2"]) > 400:
            raise FeatureError("At most 400 instances")
        _unit(_axis_vec(p["dir1"]))
        if int(p["count2"]) > 1:
            _unit(_axis_vec(p["dir2"]))
    elif p["kind"] == "circular":
        if not 2 <= int(p["count"]) <= 400:
            raise FeatureError("Circular pattern count: 2 … 400")
        _axis_spec(p["axis"])
    else:
        _plane_spec(p["plane"])


def _axis_vec(v):
    if isinstance(v, str):
        s = -1.0 if v.startswith("-") else 1.0
        k = v.strip("+-").upper()
        if k not in ("X", "Y", "Z"):
            raise FeatureError(f"Direction '{v}' must be X, Y, Z (±) or [x, y, z]")
        return np.array([s * (k == "X"), s * (k == "Y"), s * (k == "Z")], float)
    return np.asarray(v, float)


def _unit(v):
    n = float(np.linalg.norm(v))
    if n < 1e-12:
        raise FeatureError("Direction has zero length")
    return np.asarray(v, float) / n


def _axis_spec(a):
    if isinstance(a, str):
        return np.zeros(3), _unit(_axis_vec(a))
    return np.asarray(a.get("p", [0, 0, 0]), float), _unit(np.asarray(a["d"], float))


def _plane_spec(pl):
    if isinstance(pl, str):
        n = {"XY": [0, 0, 1], "XZ": [0, 1, 0], "YZ": [1, 0, 0]}.get(pl.upper())
        if n is None:
            raise FeatureError("Mirror plane: XY | XZ | YZ or {p, n}")
        return np.zeros(3), np.asarray(n, float)
    return np.asarray(pl.get("p", [0, 0, 0]), float), _unit(np.asarray(pl["n"], float))


def cam_profile(body, face_index: int, pick=None) -> dict:
    """The cam's outline: section of the body through the picked face, normal to the face's extrusion
    direction (the straight edges of a cam side face)."""
    from build123d import Face, Plane
    faces = body.faces()
    if not 0 <= face_index < len(faces):
        raise FeatureError(f"No face {face_index}")
    f = faces[face_index]
    lines = [e for e in f.edges() if e.geom_type.name == "LINE"]
    if lines:
        e = max(lines, key=lambda e_: e_.length)
        dv = e.position_at(1) - e.position_at(0)
        d = np.array([dv.X, dv.Y, dv.Z])
    elif f.geom_type.name == "PLANE":
        raise FeatureError("Pick the cam's curved side face, not a flat face")
    else:
        d = np.array([0, 0, 1.0])
    d = d / np.linalg.norm(d)
    c = f.center() if pick is None else None
    p = np.array([c.X, c.Y, c.Z]) if c is not None else np.asarray(pick, float)
    pl = Plane(origin=tuple(p), z_dir=tuple(d))
    size = body.bounding_box().diagonal * 3 + 10
    sec = body & Face.make_rect(size, size, pl)
    fs = sec.faces() if hasattr(sec, "faces") else []
    if not fs:
        raise FeatureError("Couldn't cut a cam profile through that face")
    face = max(fs, key=lambda q: q.area)
    w = face.outer_wire()
    n = max(60, min(720, int(w.length / 0.5)))
    pts = [w.position_at(i / n) for i in range(n)]
    P = [[q.X, q.Y, q.Z] for q in pts]
    return {"kind": "profile", "p": list(p), "d": list(d), "pts": P, "closed": True, "face": face_index}


def slot_ref(body, face_index: int) -> dict:
    """A slot from one of its round ends: the two parallel cylindrical faces of equal radius nearest
    each other give the end centres; d = their axis."""
    faces = body.faces()
    if not 0 <= face_index < len(faces):
        raise FeatureError(f"No face {face_index}")
    f = faces[face_index]
    if f.geom_type.name != "CYLINDER":
        raise FeatureError("Pick a round end of the slot")
    r0 = asm.ref_from_face(f, face_index)
    d0 = np.asarray(r0["d"], float)
    best, bd = None, float("inf")
    for i, g in enumerate(faces):
        if i == face_index or g.geom_type.name != "CYLINDER":
            continue
        r = asm.ref_from_face(g, i)
        if abs(r["r"] - r0["r"]) > 1e-4 or abs(abs(np.dot(np.asarray(r["d"]), d0)) - 1) > 1e-6:
            continue
        off = np.asarray(r["p"]) - np.asarray(r0["p"])
        off = off - d0 * float(off @ d0)
        dist = float(np.linalg.norm(off))
        if 1e-6 < dist < bd:
            best, bd = (r, off), dist
    if best is None:
        raise FeatureError("No matching other end found for that slot")
    p1 = np.asarray(r0["p"], float)
    p2 = p1 + best[1]
    return {"kind": "slot", "p": list(p1), "p2": list(p2), "d": list(d0), "r": r0["r"], "face": face_index}
