"""Reference geometry, loft, rib, draft, threads, sheet metal, library parts, BOM, links, explode, gear mates."""
import math
import os

import pytest

from daedalus.core.document import Document
from daedalus.core import standards as S


def box(w=40, h=30, dep=10):
    d = Document()
    d.add("sketch", {"entities": [{"type": "rect", "x": 0, "y": 0, "w": w, "h": h}]})
    d.add("extrude", {"sketch": "Sketch1", "depth": dep})
    return d


def test_loft_between_planes():
    d = Document()
    d.add("sketch", {"entities": [{"type": "rect", "x": 0, "y": 0, "w": 20, "h": 20}]})
    d.add("plane", {"mode": "offset", "base": "XY", "offset": 30})
    d.add("sketch", {"plane": "Plane1", "entities": [{"type": "circle", "x": 0, "y": 0, "r": 5}]})
    d.add("loft", {"sections": "Sketch1, Sketch2", "ruled": True})
    assert not d.result.errors
    b = d.result.bodies["Object1"]
    assert b.bounding_box().size.Z == pytest.approx(30)
    assert 30 * math.pi * 25 < b.volume < 30 * 400


def test_angled_plane_and_axis_revolve():
    d = box()
    d.add("plane", {"mode": "angle", "base": "XY", "angle": 30, "about": "x"})
    d.add("sketch", {"plane": "Plane1", "entities": [{"type": "circle", "x": 0, "y": 0, "r": 3}]})
    assert not d.result.errors
    d.add("axis", {"mode": "points", "points": [[0, 40, 0], [0, 40, 10]]})
    d.add("sketch", {"plane": "XZ", "entities": [{"type": "rect", "x": 5, "y": 5, "w": 4, "h": 4}]})
    d.add("revolve", {"sketch": "Sketch3", "axis": "Axis1", "angle": 360})
    assert not d.result.errors


def test_thread_and_hole_wizard():
    d = box(dep=20)
    d.add("hole", {"face": "top", "positions": [[0, 0]], "fastener": "M8", "fit": "tap", "depth": 15})
    th = list(d.result.threads.values())
    assert th and th[0]["major"] == 8 and th[0]["pitch"] == 1.25
    assert S.hole_size("M8", "tap") == pytest.approx(6.8)
    assert S.hole_size("M8", "normal") == pytest.approx(9.0)


def test_draft_sides():
    d = box()
    v0 = d.result.bodies["Object1"].volume
    d.add("draft", {"faces": "sides", "angle": 3, "neutral": "bottom"})
    assert not d.result.errors
    assert d.result.bodies["Object1"].volume < v0


def test_sheet_metal_flat_size():
    d = Document()
    d.add("sketch", {"entities": [{"type": "rect", "x": 0, "y": 0, "w": 100, "h": 60}]})
    t, R, K = 2.0, 2.0, 0.44
    d.add("sheet_metal", {"sketch": "Sketch1", "thickness": t, "bend_radius": R, "k_factor": K,
                          "flanges": [{"edge": "+X", "length": 30, "angle": 90}]})
    assert not d.result.errors
    fl = d.result.flats["SheetMetal1"]
    ba = math.pi / 2 * (R + K * t)
    # flange length measured to the outside: flat = base + flange − 2 × setback + BA (setback = R + t)
    assert fl["size"][1] == pytest.approx(60)
    assert len(fl["bends"]) == 1
    assert 100 + 30 - 2 * (R + t) + ba - 3 < fl["size"][0] < 100 + 30 + ba + 1


def test_library_screw_bom_and_explode(tmp_path):
    d = box(dep=10)
    d.add("hole", {"face": "top", "positions": [[10, 5]], "fastener": "M8", "fit": "normal"})
    r = d.insert_part("socket head cap screw", "M8", 25, target={"obj": "Object1", "face": {"axis_near": [10, 5, 5]}},
                      washer=True, nut=True)
    assert len(r["objects"]) == 3 and len(r["mates"]) >= 4
    assert not d.interference()
    rows = d.bom()
    assert {row["part"] for row in rows} >= {"Object1"}
    assert any("M8" in row["part"] for row in rows)
    d.export_bom(tmp_path / "bom.xlsx")
    import openpyxl
    wb = openpyxl.load_workbook(tmp_path / "bom.xlsx")
    assert "BOM" in wb.sheetnames
    off = d.auto_explode(checkpoint=False)
    assert off and d.exploded
    poses = d.display_poses()
    screw = next(n for n in r["objects"] if n.startswith("SHCS"))
    assert poses[screw] != d.obj(screw).pose


def test_linked_part_updates(tmp_path):
    src = box()
    src.save(tmp_path / "part.daedalus")
    d = Document()
    n = d.insert_part_file(tmp_path / "part.daedalus", link=True)
    v0 = d.result.bodies[n].volume
    src.edit("Extrude1", {"depth": 20})
    src.save(tmp_path / "part.daedalus")
    os.utime(tmp_path / "part.daedalus", None)
    import time
    time.sleep(0.01)
    st = os.stat(tmp_path / "part.daedalus")
    os.utime(tmp_path / "part.daedalus", (st.st_atime, st.st_mtime + 5))
    assert d.update_links()
    d.rebuild()
    assert d.result.bodies[n].volume == pytest.approx(2 * v0)


def test_gear_mate_ratio():
    d = Document()
    d.add("sketch", {"entities": [{"type": "circle", "x": 0, "y": 0, "r": 10}]})
    d.add("extrude", {"sketch": "Sketch1", "depth": 5})
    d.add_object("G2", pose=[40, 0, 0, 0, 0, 0])
    d.add("sketch", {"entities": [{"type": "circle", "x": 0, "y": 0, "r": 20}]}, obj="G2")
    d.add("extrude", {"sketch": "Sketch2", "depth": 5}, obj="G2")
    d.obj("Object1").fixed = False
    d.add_object("Ground", pose=[0, 0, -20, 0, 0, 0])
    m = d.add_mate("gear", {"obj": "Object1", "ref": {"axis_near": [10, 0, 2.5]}},
                   {"obj": "G2", "ref": {"axis_near": [60, 0, 2.5]}}, value=2.0)
    assert m.type == "gear"
    from daedalus.core import assembly as asm
    assert "gear" in asm.MATE_TYPES and "rack" in asm.MATE_TYPES


def test_sketch_fillet_offset_mirror():
    from daedalus.core import sketch_ops as ops
    g = {"points": {"p1": [0, 0], "p2": [20, 0], "p3": [20, 10], "p4": [0, 10]},
         "curves": [{"id": "l1", "type": "line", "a": "p1", "b": "p2"}, {"id": "l2", "type": "line", "a": "p2", "b": "p3"},
                    {"id": "l3", "type": "line", "a": "p3", "b": "p4"}, {"id": "l4", "type": "line", "a": "p4", "b": "p1"}],
         "constraints": []}
    f = ops.fillet(g, "p2", 3.0)
    assert any(c["type"] == "arc" for c in f["curves"])
    o = ops.offset(g, ["l1", "l2", "l3", "l4"], 2.0)
    assert len(o["curves"]) >= 8
    new = sorted(tuple(v) for k, v in o["points"].items() if k not in g["points"])
    assert new == [(2.0, 2.0), (2.0, 8.0), (18.0, 2.0), (18.0, 8.0)]       # offset 2 mm inside
